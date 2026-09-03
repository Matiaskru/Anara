"""Leitura da planilha Anara pra importação manual (Etapa 2).

Reaproveita o mesmo padrão do gerar_cotacao.py: abre com openpyxl
data_only=True e lê os valores JÁ CALCULADOS pelo Excel — não reimplementa
a cadeia de nacionalização (01_CUSTOS -> 02_RESUMO_VISUAL ->
03_Precificação Anara), só consome o resultado final dela.
"""
import json
import warnings
from dataclasses import dataclass, field
from typing import List, Optional

import openpyxl

warnings.filterwarnings("ignore")

SHEET_PRECIFICACAO = "03_Precificação Anara"
SHEET_RESUMO = "02_RESUMO_VISUAL"
SHEET_PREMISSAS = "05_Premissas"
SHEET_DIFAL = "06_DIFAL_Estados"
SHEET_CENARIOS = "07_Cenarios_Fiscais"
SHEET_CUSTOS = "01_CUSTOS"

# bloco central de premissas de nacionalização (05_Premissas, reforma ago/2026)
CEL_CAMBIO = "C114"
CEL_FRETE_USD_KG = "C117"
CEL_OUTRAS_DESP = "C118"

COMISSAO_LINHA_INICIO = 91
COMISSAO_LINHA_FIM = 96


@dataclass
class ProdutoImportado:
    sku_key: str
    categoria: Optional[str]
    nome: str
    especificacao: Optional[str]
    custo_unitario: float
    preco_base: float
    # memória de nacionalização (interna, nunca vai pro PDF do cliente)
    ncm: Optional[str] = None
    ii_aplicado: Optional[float] = None
    peso_kg: Optional[float] = None
    peso_fonte: Optional[str] = None
    peso_tipo: Optional[str] = None
    preco_ktc_usd: Optional[float] = None
    frete_usd_un: Optional[float] = None
    custo_net_usd: Optional[float] = None
    cotacao_origem: Optional[str] = None


@dataclass
class ImportResult:
    produtos: List[ProdutoImportado] = field(default_factory=list)
    icms_pct: float = 0.0
    pis_cofins_pct: float = 0.0
    encargo_financeiro_pct: float = 0.0
    comissao_tabela: List[tuple] = field(default_factory=list)
    origem_uf: str = "SC"
    icms_por_estado: dict = field(default_factory=dict)  # {"São Paulo": 0.1707, ...} — Carga Final
    cenarios_fiscais: List[dict] = field(default_factory=list)  # tabela central Origem×Destino×Contribuinte
    # premissas de nacionalização (bloco central 05_Premissas!B113+, reforma ago/2026)
    cambio_usd_brl: Optional[float] = None
    frete_usd_kg: Optional[float] = None
    outras_desp_usd_un: Optional[float] = None
    aviso: Optional[str] = None  # preenchido se o Excel não tinha valores calculados em cache


def ler_excel(path: str) -> ImportResult:
    wb = openpyxl.load_workbook(path, data_only=True)

    if SHEET_PRECIFICACAO not in wb.sheetnames:
        return ImportResult(aviso=f"Aba '{SHEET_PRECIFICACAO}' não encontrada no arquivo.")

    ws = wb[SHEET_PRECIFICACAO]
    ws_custos = wb[SHEET_CUSTOS] if SHEET_CUSTOS in wb.sheetnames else None
    produtos = []
    valores_em_branco = 0
    linhas_com_produto = 0

    def memoria_nacionalizacao(linha_precificacao: int, nome_esperado: str) -> dict:
        """01_CUSTOS linha N corresponde a 03_Precificação linha N+1. Só usa a
        memória se o nome do produto bater — senão devolve vazio (seguro)."""
        if not ws_custos:
            return {}
        rc = linha_precificacao - 1
        if ws_custos[f"D{rc}"].value != nome_esperado:
            return {}

        def num(ref):
            v = ws_custos[ref].value
            return float(v) if isinstance(v, (int, float)) else None

        return {
            "ncm": ws_custos[f"P{rc}"].value,
            "ii_aplicado": num(f"AC{rc}"),
            "peso_kg": num(f"BG{rc}"),
            "peso_fonte": ws_custos[f"BI{rc}"].value,
            "peso_tipo": ws_custos[f"BH{rc}"].value,
            "preco_ktc_usd": num(f"R{rc}"),
            "frete_usd_un": num(f"BJ{rc}"),
            "custo_net_usd": num(f"BN{rc}"),
            "cotacao_origem": ws_custos[f"B{rc}"].value,
        }

    for r in range(3, ws.max_row + 1):
        nome = ws[f"B{r}"].value
        if not nome:
            continue
        linhas_com_produto += 1
        custo = ws[f"D{r}"].value
        preco = ws[f"F{r}"].value
        sku_key = ws[f"Y{r}"].value or f"{nome}  ·  {ws[f'C{r}'].value or ''}"
        if custo is None or preco is None:
            valores_em_branco += 1
            continue
        produtos.append(ProdutoImportado(
            sku_key=sku_key,
            categoria=ws[f"A{r}"].value,
            nome=nome,
            especificacao=ws[f"C{r}"].value,
            custo_unitario=float(custo),
            preco_base=float(preco),
            **memoria_nacionalizacao(r, nome),
        ))

    aviso = None
    if linhas_com_produto and valores_em_branco == linhas_com_produto:
        aviso = (
            "Os preços/custos estão em branco no Excel (arquivo foi salvo por um script, não "
            "pelo Excel). Abra o arquivo no Excel, aperte Cmd+S, e importe de novo."
        )

    icms_pct = pis_cofins_pct = encargo_pct = 0.0
    if SHEET_RESUMO in wb.sheetnames:
        wr = wb[SHEET_RESUMO]
        icms_pct = wr["P2"].value or 0.0
        pis_cofins_pct = wr["P3"].value or 0.0
        encargo_pct = wr["P4"].value or 0.0

    comissao_tabela = []
    cambio = frete_kg = outras_desp = None
    if SHEET_PREMISSAS in wb.sheetnames:
        wp = wb[SHEET_PREMISSAS]
        for r in range(COMISSAO_LINHA_INICIO, COMISSAO_LINHA_FIM + 1):
            markup_min = wp[f"B{r}"].value
            comissao = wp[f"C{r}"].value
            if isinstance(markup_min, (int, float)) and isinstance(comissao, (int, float)):
                comissao_tabela.append((float(markup_min), float(comissao)))

        def prem(ref):
            v = wp[ref].value
            return float(v) if isinstance(v, (int, float)) else None

        cambio = prem(CEL_CAMBIO)
        frete_kg = prem(CEL_FRETE_USD_KG)
        outras_desp = prem(CEL_OUTRAS_DESP)

    icms_por_estado = {}
    if SHEET_DIFAL in wb.sheetnames:
        wd = wb[SHEET_DIFAL]
        for r in range(2, wd.max_row + 1):
            estado = wd[f"A{r}"].value
            carga_final = wd[f"H{r}"].value
            if estado and isinstance(carga_final, (int, float)):
                icms_por_estado[estado] = float(carga_final)

    cenarios_fiscais = []
    if SHEET_CENARIOS in wb.sheetnames:
        wc = wb[SHEET_CENARIOS]
        for r in range(2, wc.max_row + 1):
            origem = wc[f"A{r}"].value
            destino = wc[f"B{r}"].value
            contribuinte_txt = wc[f"C{r}"].value
            icms_venda = wc[f"D{r}"].value
            regra = wc[f"E{r}"].value
            if origem and destino and isinstance(icms_venda, (int, float)):
                cenarios_fiscais.append({
                    "origem": origem, "destino": destino,
                    "contribuinte": str(contribuinte_txt).strip().upper() == "SIM",
                    "icms_venda": float(icms_venda), "regra": regra or "",
                })

    return ImportResult(
        produtos=produtos, icms_pct=icms_pct, pis_cofins_pct=pis_cofins_pct,
        encargo_financeiro_pct=encargo_pct, comissao_tabela=comissao_tabela,
        icms_por_estado=icms_por_estado, cenarios_fiscais=cenarios_fiscais,
        cambio_usd_brl=cambio, frete_usd_kg=frete_kg, outras_desp_usd_un=outras_desp,
        aviso=aviso,
    )


@dataclass
class DiffProduto:
    sku_key: str
    nome: str
    especificacao: Optional[str]
    tipo: str  # "novo" | "removido" | "alterado" | "igual"
    preco_base_antigo: Optional[float] = None
    preco_base_novo: Optional[float] = None
    custo_antigo: Optional[float] = None
    custo_novo: Optional[float] = None


def montar_diff(produtos_atuais: dict, resultado: ImportResult) -> List[DiffProduto]:
    """`produtos_atuais`: dict sku_key -> objeto com .preco_base/.custo_unitario/.ativo."""
    diffs = []
    novos_keys = set()
    for p in resultado.produtos:
        novos_keys.add(p.sku_key)
        atual = produtos_atuais.get(p.sku_key)
        if atual is None:
            diffs.append(DiffProduto(p.sku_key, p.nome, p.especificacao, "novo",
                                      preco_base_novo=p.preco_base, custo_novo=p.custo_unitario))
        else:
            mudou = (abs((atual.preco_base or 0) - p.preco_base) > 0.005
                     or abs((atual.custo_unitario or 0) - p.custo_unitario) > 0.005)
            diffs.append(DiffProduto(
                p.sku_key, p.nome, p.especificacao, "alterado" if mudou else "igual",
                preco_base_antigo=atual.preco_base, preco_base_novo=p.preco_base,
                custo_antigo=atual.custo_unitario, custo_novo=p.custo_unitario,
            ))
    for sku_key, atual in produtos_atuais.items():
        if sku_key not in novos_keys and atual.ativo:
            diffs.append(DiffProduto(sku_key, atual.nome, atual.especificacao, "removido",
                                      preco_base_antigo=atual.preco_base,
                                      custo_antigo=atual.custo_unitario))
    return diffs


def comissao_tabela_to_json(tabela: List[tuple]) -> str:
    return json.dumps(tabela)


def comissao_tabela_from_json(s: str) -> List[tuple]:
    return [tuple(x) for x in json.loads(s)]
