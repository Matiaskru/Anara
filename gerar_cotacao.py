#!/usr/bin/env python3
"""Gera um PDF de proposta comercial Anara a partir da aba COTAÇÃO do Excel."""
import datetime
import os
import subprocess
import sys
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Table, TableStyle,
)

# ---------------------------------------------------------------------------
# Config — ajuste o caminho do Excel aqui se você mover o arquivo de lugar
# ---------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.expanduser(
    "~/Downloads/Sistema de preços Anara - II Egito corrigido (atualizado).xlsx"
)
OUTPUT_DIR = os.path.expanduser("~/Desktop")
LOGO_PATH = os.path.join(BASE_DIR, "assets", "logo", "anara_mark_copper.png")
FONTS_DIR = os.path.join(BASE_DIR, "assets", "fonts")

# ---------------------------------------------------------------------------
# Marca — cores e fontes (Proposta 1 do branding, Mar/2026)
# ---------------------------------------------------------------------------
MIDNIGHT_BLUE = colors.HexColor("#30354F")
COPPER = colors.HexColor("#C49281")
SAND_BEIGE = colors.HexColor("#D6C8BE")
COTTON = colors.HexColor("#F9F6F2")
COOL_GRAY = colors.HexColor("#6B6E7C")
ROW_ALT = colors.HexColor("#F1EBE6")

pdfmetrics.registerFont(TTFont("Didot", os.path.join(FONTS_DIR, "Didot.ttf")))
pdfmetrics.registerFont(TTFont("Didot-Bold", os.path.join(FONTS_DIR, "Didot-Bold.ttf")))
pdfmetrics.registerFont(TTFont("Futura", os.path.join(FONTS_DIR, "Futura-Medium.ttf")))
pdfmetrics.registerFont(TTFont("Futura-Bold", os.path.join(FONTS_DIR, "Futura-Bold.ttf")))

PAGE_W, PAGE_H = A4
MARGIN = 18 * mm
HEADER_H = 32 * mm
FOOTER_H = 16 * mm


def brl(v):
    if v in (None, ""):
        return ""
    try:
        v = float(v)
    except (TypeError, ValueError):
        return str(v)
    s = f"{v:,.2f}"
    s = s.replace(",", "§").replace(".", ",").replace("§", ".")
    return f"R$ {s}"


def pct(v):
    if v in (None, ""):
        return "-"
    try:
        v = float(v)
    except (TypeError, ValueError):
        return str(v)
    return f"{v*100:.0f}%" if v <= 1 else f"{v:.0f}%"


# ---------------------------------------------------------------------------
# Leitura do Excel
# ---------------------------------------------------------------------------
def load_quote(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["COTAÇÃO"]

    def cell(ref):
        return ws[ref].value

    header = {
        "cliente": cell("B4"),
        "cnpj": cell("B5"),
        "cidade": cell("B6"),
        "telefone": cell("B7"),
        "email": cell("B8"),
        "obs": cell("B9"),
        "data": cell("E4"),
        "validade": cell("E5"),
        "vendedor": cell("E6"),
        "condicao_pagamento": cell("E7") or "30",  # única pra cotação inteira
        "frete": cell("E8"),
    }

    items = []
    for r in range(12, 52):
        produto = ws[f"B{r}"].value
        qtd = ws[f"D{r}"].value
        if not produto or not qtd:
            continue
        items.append({
            "n": ws[f"A{r}"].value,
            "produto": produto,
            "spec": ws[f"C{r}"].value,
            "qtd": ws[f"D{r}"].value,
            "preco_unit": ws[f"E{r}"].value,
            "desc": ws[f"F{r}"].value,
            "preco_final": ws[f"G{r}"].value,
            "total": ws[f"H{r}"].value,
        })

    totals = {
        "subtotal": cell("H53"),
        "frete": cell("H54"),
        "total_geral": cell("H55"),
        "total_itens": cell("H56"),
    }
    return header, items, totals


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def draw_header(canvas, doc, header):
    canvas.saveState()
    canvas.setFillColor(MIDNIGHT_BLUE)
    canvas.rect(0, PAGE_H - HEADER_H, PAGE_W, HEADER_H, stroke=0, fill=1)

    logo_size = 16 * mm
    logo_x = MARGIN
    logo_y = PAGE_H - HEADER_H + (HEADER_H - logo_size) / 2
    canvas.drawImage(LOGO_PATH, logo_x, logo_y, width=logo_size, height=logo_size,
                      mask="auto")

    text_x = logo_x + logo_size + 5 * mm
    canvas.setFillColor(COPPER)
    canvas.setFont("Didot-Bold", 22)
    canvas.drawString(text_x, PAGE_H - HEADER_H + 17.5 * mm, "A N A R A")
    canvas.setFillColor(COTTON)
    canvas.setFont("Futura", 7.5)
    canvas.drawString(text_x, PAGE_H - HEADER_H + 12.5 * mm,
                       "A   H I G H E R   S T A N D A R D   O F   C O M F O R T")

    canvas.setFillColor(COTTON)
    canvas.setFont("Futura-Bold", 12)
    canvas.drawRightString(PAGE_W - MARGIN, PAGE_H - HEADER_H + 19 * mm,
                            "PROPOSTA COMERCIAL")
    canvas.setFont("Futura", 8.5)
    data_txt = header.get("data") or ""
    validade_txt = header.get("validade") or ""
    canvas.drawRightString(PAGE_W - MARGIN, PAGE_H - HEADER_H + 13.5 * mm,
                            f"Data: {data_txt}")
    canvas.drawRightString(PAGE_W - MARGIN, PAGE_H - HEADER_H + 9.5 * mm,
                            f"Validade: {validade_txt}")
    numero = header.get("numero")
    if numero:
        canvas.setFont("Futura-Bold", 8.5)
        canvas.drawRightString(PAGE_W - MARGIN, PAGE_H - HEADER_H + 5.5 * mm, f"Nº {numero}")
    canvas.restoreState()


def draw_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(COPPER)
    canvas.setLineWidth(0.75)
    canvas.line(MARGIN, FOOTER_H, PAGE_W - MARGIN, FOOTER_H)
    canvas.setFillColor(COOL_GRAY)
    canvas.setFont("Futura", 6.3)
    # O cliente vê preço, nunca a composição interna: nada de ICMS, PIS/COFINS, comissão,
    # markup, margem, custo ou câmbio discriminados aqui.
    line1 = ("Preços em reais, com impostos inclusos, válidos para as condições comerciais desta "
             "proposta. Alteração de prazo, destino ou condição de pagamento pode alterar os preços.")
    line2 = "Anara — enxoval hoteleiro."
    canvas.drawString(MARGIN, FOOTER_H - 4.5 * mm, line1)
    canvas.drawString(MARGIN, FOOTER_H - 8.5 * mm, line2)
    canvas.setFont("Futura", 7)
    canvas.drawRightString(PAGE_W - MARGIN, FOOTER_H - 8.5 * mm, f"Página {doc.page}")
    canvas.restoreState()


def build_pdf(out_path, header, items, totals):
    doc = BaseDocTemplate(
        out_path, pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=HEADER_H + 6 * mm, bottomMargin=FOOTER_H + 4 * mm,
    )
    frame = Frame(MARGIN, FOOTER_H + 4 * mm, PAGE_W - 2 * MARGIN,
                   PAGE_H - HEADER_H - FOOTER_H - 10 * mm, id="body")

    def on_page(canvas, doc_):
        draw_header(canvas, doc_, header)
        draw_footer(canvas, doc_)

    doc.addPageTemplates([PageTemplate(id="anara", frames=[frame], onPage=on_page)])

    label_style = ParagraphStyle("label", fontName="Futura", fontSize=7.5,
                                  textColor=COOL_GRAY, leading=9)
    value_style = ParagraphStyle("value", fontName="Futura-Bold", fontSize=9.5,
                                  textColor=MIDNIGHT_BLUE, leading=12)

    def field(label, value):
        return Paragraph(
            f'<font name="Futura" size="7.5" color="#6B6E7C">{label}</font><br/>'
            f'<font name="Futura-Bold" size="9.5" color="#30354F">{value or "—"}</font>',
            value_style,
        )

    story = []

    client_rows = [
        [field("CLIENTE", header.get("cliente")), field("VENDEDOR", header.get("vendedor"))],
        [field("CNPJ/CPF", header.get("cnpj")), field("CONDIÇÃO DE PAGAMENTO", header.get("condicao_pagamento"))],
        [field("CIDADE/UF", header.get("cidade")), field("FRETE", header.get("frete"))],
        [field("CONTATO", header.get("contato")), field("PRAZO DE ENTREGA", header.get("prazo_entrega"))],
        [field("DEPARTAMENTO", header.get("departamento")), field("LOCAL DE ENTREGA", header.get("local_entrega"))],
        [field("TELEFONE", header.get("telefone")), field("E-MAIL", header.get("email"))],
    ]
    client_table = Table(client_rows, colWidths=[85 * mm, 79 * mm])
    client_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("LINEBELOW", (0, -1), (-1, -1), 0.5, SAND_BEIGE),
    ]))
    story.append(client_table)
    story.append(Spacer(1, 3 * mm))

    if header.get("obs"):
        story.append(Paragraph(
            f'<font name="Futura" size="7.5" color="#6B6E7C">OBS</font><br/>'
            f'<font name="Futura" size="9" color="#30354F">{header["obs"]}</font>',
            value_style))
        story.append(Spacer(1, 3 * mm))

    story.append(Spacer(1, 2 * mm))

    col_widths = [7*mm, 62*mm, 13*mm, 23*mm, 12*mm, 23*mm, 24*mm]
    head_style = ParagraphStyle("thead", fontName="Futura-Bold", fontSize=7.3,
                                 textColor=COTTON, leading=9, alignment=TA_LEFT)
    cell_style = ParagraphStyle("tcell", fontName="Futura", fontSize=8, leading=10,
                                 textColor=MIDNIGHT_BLUE)
    cell_style_r = ParagraphStyle("tcell_r", parent=cell_style, alignment=TA_RIGHT)
    spec_style = ParagraphStyle("spec", fontName="Futura", fontSize=6.8, leading=8.4,
                                 textColor=COOL_GRAY)

    header_row = [
        Paragraph("#", head_style), Paragraph("PRODUTO", head_style),
        Paragraph("QTD", head_style),
        Paragraph("PREÇO<br/>UNIT", head_style), Paragraph("DESC", head_style),
        Paragraph("PREÇO<br/>C/DESC", head_style), Paragraph("TOTAL", head_style),
    ]
    data = [header_row]
    for it in items:
        prod_cell = Paragraph(
            f'<font name="Futura-Bold" size="8" color="#30354F">{it["produto"]}</font>'
            + (f'<br/><font name="Futura" size="6.8" color="#6B6E7C">{it["spec"]}</font>'
               if it.get("spec") else ""),
            cell_style,
        )
        data.append([
            Paragraph(str(it["n"] or ""), cell_style),
            prod_cell,
            Paragraph(str(it["qtd"] or ""), cell_style_r),
            Paragraph(brl(it["preco_unit"]), cell_style_r),
            Paragraph(pct(it["desc"]), cell_style_r),
            Paragraph(brl(it["preco_final"]), cell_style_r),
            Paragraph(brl(it["total"]), cell_style_r),
        ])

    items_table = Table(data, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), MIDNIGHT_BLUE),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        ("TOPPADDING", (0, 1), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 3.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, MIDNIGHT_BLUE),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, SAND_BEIGE),
    ]
    for i in range(1, len(data)):
        if i % 2 == 0:
            style.append(("BACKGROUND", (0, i), (-1, i), ROW_ALT))
    items_table.setStyle(TableStyle(style))
    story.append(items_table)
    story.append(Spacer(1, 5 * mm))

    tot_style_label = ParagraphStyle("totlabel", fontName="Futura", fontSize=9,
                                      textColor=COOL_GRAY, alignment=TA_RIGHT)
    tot_style_val = ParagraphStyle("totval", fontName="Futura-Bold", fontSize=9,
                                    textColor=MIDNIGHT_BLUE, alignment=TA_RIGHT)
    tot_style_val_big = ParagraphStyle("totvalbig", fontName="Didot-Bold", fontSize=13.5,
                                        textColor=COTTON, alignment=TA_RIGHT)
    tot_style_label_big = ParagraphStyle("totlabelbig", fontName="Futura-Bold", fontSize=10,
                                          textColor=COTTON, alignment=TA_RIGHT)

    totals_data = [
        [Paragraph("SUBTOTAL", tot_style_label), Paragraph(brl(totals["subtotal"]), tot_style_val)],
        [Paragraph("FRETE", tot_style_label), Paragraph(brl(totals["frete"]) or "A combinar", tot_style_val)],
        [Paragraph("TOTAL GERAL", tot_style_label_big), Paragraph(brl(totals["total_geral"]), tot_style_val_big)],
    ]
    totals_table = Table(totals_data, colWidths=[36 * mm, 46 * mm])
    totals_table.setStyle(TableStyle([
        ("TOPPADDING", (0, 0), (-1, 1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, 1), 2),
        ("BACKGROUND", (0, 2), (-1, 2), MIDNIGHT_BLUE),
        ("TOPPADDING", (0, 2), (-1, 2), 5),
        ("BOTTOMPADDING", (0, 2), (-1, 2), 5),
        ("RIGHTPADDING", (0, 2), (-1, 2), 6),
        ("LEFTPADDING", (0, 2), (-1, 2), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    wrapper = Table([[Spacer(1, 1), totals_table]], colWidths=[92 * mm, 82 * mm])
    wrapper.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM")]))
    story.append(wrapper)

    termos = (header.get("termos") or "").strip()
    if termos:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(
            '<font name="Futura-Bold" size="8" color="#30354F">TERMOS E CONDIÇÕES</font>',
            value_style))
        story.append(Spacer(1, 1.5 * mm))
        termos_style = ParagraphStyle("termos", fontName="Futura", fontSize=7.6, leading=10.4,
                                       textColor=COOL_GRAY)
        for paragrafo in [t for t in termos.split("\n") if t.strip()]:
            story.append(Paragraph(paragrafo.strip(), termos_style))

    if header.get("mostrar_aceite"):
        story.append(Spacer(1, 7 * mm))
        aceite_label = ParagraphStyle("aceitelabel", fontName="Futura", fontSize=7,
                                       textColor=COOL_GRAY, leading=9)
        linha_assinatura = Table(
            [[Paragraph("DE ACORDO — NOME E CARGO", aceite_label),
              Paragraph("DATA", aceite_label)]],
            colWidths=[118 * mm, 56 * mm])
        linha_assinatura.setStyle(TableStyle([
            ("LINEABOVE", (0, 0), (-1, 0), 0.5, SAND_BEIGE),
            ("TOPPADDING", (0, 0), (-1, 0), 3),
        ]))
        story.append(Spacer(1, 9 * mm))
        story.append(linha_assinatura)

    doc.build(story)


def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"Não encontrei o Excel em: {EXCEL_PATH}", file=sys.stderr)
        sys.exit(1)

    header, items, totals = load_quote(EXCEL_PATH)
    if not items:
        print("Nenhum item preenchido na aba COTAÇÃO (linhas 12-51, coluna B). Nada a gerar.",
              file=sys.stderr)
        sys.exit(1)

    if totals.get("total_geral") is None or all(it.get("total") is None for it in items):
        msg = (
            "Os preços/totais estão em branco no Excel.\n\n"
            "Isso acontece quando o arquivo foi salvo por um script (não pelo Excel), então as "
            "fórmulas ainda não têm valor calculado guardado.\n\n"
            "Abra o arquivo no Excel (ele recalcula sozinho ao abrir), aperte Cmd+S, e tente "
            "gerar de novo. Depois disso o gerador sempre vai pegar os valores certos, direto "
            "do que estiver salvo."
        )
        print(msg, file=sys.stderr)
        subprocess.run(["open", EXCEL_PATH], check=False)
        sys.exit(1)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    cliente = (header.get("cliente") or "Cliente").strip()
    safe_cliente = "".join(c if c.isalnum() or c in " -_" else "" for c in cliente).strip() or "Cliente"
    stamp = datetime.datetime.now().strftime("%d-%m-%Y_%H%M")
    out_path = os.path.join(OUTPUT_DIR, f"Cotação Anara - {safe_cliente} - {stamp}.pdf")

    build_pdf(out_path, header, items, totals)
    print(f"PDF gerado: {out_path}")

    try:
        subprocess.run(["open", out_path], check=False)
    except Exception:
        pass


if __name__ == "__main__":
    main()
