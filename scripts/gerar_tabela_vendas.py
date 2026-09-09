#!/usr/bin/env python3
"""Tabela de preços para as vendedoras — a versão que pode sair da sala.

Arquivo **separado** de `Lista de Precos Anara.xlsx`, e essa separação é o ponto: a lista
interna tem custo NET, margem alvo, margem real, comissão, lucro e imposto por linha. Nada
disso pode chegar a quem vende. A regra do projeto é explícita — vendedor não vê custo, EXW,
custo NET, margem interna, markup, impostos detalhados nem premissas — e mandar o arquivo
interno "pedindo para não olhar a aba INTERNO" não é controle de acesso, é torcida.

Aqui não existe nenhuma dessas colunas. O que existe é o que a vendedora precisa para fechar
uma proposta sozinha: escolher o destino, dizer se o cliente é contribuinte, escolher o prazo
de pagamento, digitar as quantidades e ler o total.

O preço é **posto fábrica**: frete não está incluído. Não é omissão — a tarifa da TRANSAL é
por tonelada com mínimo por embarque, e só a Daune e a Decor nem peso cadastrado têm. Somar um
frete por peça aqui seria inventar precisão que o dado não tem.

Rodar:  python3 scripts/gerar_tabela_vendas.py
"""
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.expanduser("~/Anara-Cotacao"))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import create_engine
from sqlmodel import Session

from scripts.gerar_lista_precos import calcular, instalar_derivacao

DB = os.path.expanduser("~/Anara-Cotacao/data/anara.db")
SAIDA = os.path.expanduser("~/Downloads/Tabela de Precos ANARA - Vendas.xlsx")

MIDNIGHT, COPPER, SAND = "FF30354F", "FFC49281", "FFD6C8BE"
COTTON, CINZA = "FFF9F6F2", "FF6B6E7C"
AZUL, BEGE, VERMELHO, VERDE = "FFD8E4F5", "FFF4EFEA", "FFF6E0DC", "FFE3EFE7"
F = "Calibri"
BRL = '"R$ "#,##0.00'

fill_capa = PatternFill("solid", fgColor=MIDNIGHT)
fill_cab = PatternFill("solid", fgColor=MIDNIGHT)
fill_azul = PatternFill("solid", fgColor=AZUL)
fill_bege = PatternFill("solid", fgColor=BEGE)
fill_verm = PatternFill("solid", fgColor=VERMELHO)
fill_verde = PatternFill("solid", fgColor=VERDE)
lado = Side(style="thin", color=SAND)
borda = Border(left=lado, right=lado, top=lado, bottom=lado)
grosso = Side(style="medium", color=COPPER)
borda_entrada = Border(left=grosso, right=grosso, top=grosso, bottom=grosso)


def cabecalho(ws, linha, titulos):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(name=F, bold=True, color=COTTON, size=10)
        c.fill = fill_cab
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
    ws.row_dimensions[linha].height = 28


def nota_do_cenario(estado, contribuinte, origem):
    """O que a vendedora precisa dizer ao cliente sobre imposto — em português, sem alíquota.

    A alíquota é "imposto detalhado" e não vai para a vendedora. Mas *quem paga o DIFAL* é
    informação comercial: numa venda interestadual a contribuinte, o cliente recolhe esse
    imposto por fora, e uma proposta que não avisa isso vira discussão depois.
    """
    if estado == origem:
        return "Venda dentro do estado. Nenhum imposto adicional por fora deste preço."
    if contribuinte:
        return ("Venda interestadual para CONTRIBUINTE: além deste preço, o cliente recolhe "
                "o DIFAL do estado dele. Diga isso na proposta.")
    return ("Venda interestadual para NÃO CONTRIBUINTE: todo o imposto já está dentro deste "
            "preço. O cliente não recolhe nada por fora.")


def escrever(grade, ctx):
    wb = Workbook()
    cods = [c.codigo for c in ctx["condicoes"]]
    origem = ctx["origem_nome"]

    # ------------------------------------------------------------- PRECOS
    pr = wb.active
    pr.title = "PRECOS"
    cabecalho(pr, 1, ["CHAVE", "Produto", "Destino", "Contribuinte"] + cods)
    linha = 2
    vistos, ordem_skus = set(), []
    for (sku, est, contrib), precos in sorted(grade.items()):
        pr.cell(row=linha, column=1, value=f"{sku}|{est}|{'SIM' if contrib else 'NAO'}")
        pr.cell(row=linha, column=2, value=sku)
        pr.cell(row=linha, column=3, value=est)
        pr.cell(row=linha, column=4, value="SIM" if contrib else "NAO")
        for i, cod in enumerate(cods, start=5):
            pr.cell(row=linha, column=i, value=precos.get(cod)).number_format = BRL
        if sku not in vistos:
            vistos.add(sku)
            ordem_skus.append(sku)
        linha += 1
    fim = linha - 1
    pr.sheet_state = "hidden"

    # ------------------------------------------------------------- LISTAS
    lst = wb.create_sheet("LISTAS")
    estados = ctx["estados"]
    for i, e in enumerate(estados, start=1):
        lst.cell(row=i, column=1, value=e.estado)
    lst.cell(row=1, column=2, value="SIM")
    lst.cell(row=2, column=2, value="NAO")
    for i, cod in enumerate(cods, start=1):
        lst.cell(row=i, column=3, value=cod)
    i = 1
    for e in estados:
        for contrib in (True, False):
            lst.cell(row=i, column=5, value=f"{e.estado}|{'SIM' if contrib else 'NAO'}")
            lst.cell(row=i, column=6, value=nota_do_cenario(e.estado, contrib, origem))
            i += 1
    fim_cen = i - 1
    lst.sheet_state = "hidden"

    # -------------------------------------------------------------- COTAR
    ws = wb.create_sheet("COTAR", 0)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [3, 46, 40, 15, 12, 17, 3]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["H"].hidden = True

    ws.merge_cells("B2:F2")
    t = ws["B2"]
    t.value = "ANARA — TABELA DE PREÇOS"
    t.font = Font(name=F, bold=True, size=18, color=COTTON)
    t.fill = fill_capa
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 40

    ws.merge_cells("B3:F3")
    s = ws["B3"]
    s.value = ("1) Escolha destino, tipo de cliente e prazo.   2) Digite as quantidades na "
               "coluna QUANTIDADE.   3) O total sai sozinho.")
    s.font = Font(name=F, size=10, italic=True, color=CINZA)
    s.alignment = Alignment(horizontal="center")

    for rot, cel, texto, padrao in (("B5", "C5", "Estado de destino", origem),
                                    ("B6", "C6", "Cliente é contribuinte de ICMS?", "SIM"),
                                    ("B7", "C7", "Prazo de pagamento", cods[0])):
        r = ws[rot]
        r.value = texto
        r.font = Font(name=F, bold=True, size=11, color=MIDNIGHT)
        r.alignment = Alignment(horizontal="right", vertical="center")
        c = ws[cel]
        c.value = padrao
        c.font = Font(name=F, bold=True, size=12, color=MIDNIGHT)
        c.fill = fill_azul
        c.border = borda_entrada
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[int(rot[1:])].height = 24

    for cel, formula in (("C5", f"LISTAS!$A$1:$A${len(estados)}"),
                         ("C6", "LISTAS!$B$1:$B$2"),
                         ("C7", f"LISTAS!$C$1:$C${len(cods)}")):
        dv = DataValidation(type="list", formula1=f"={formula}", allow_blank=False,
                            showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(ws[cel])

    # Total da cotação, ao lado das escolhas: quem cota quer ver o número sem rolar a lista.
    rot = ws["E5"]
    rot.value = "TOTAL DA COTAÇÃO"
    rot.font = Font(name=F, bold=True, size=10, color=COTTON)
    rot.fill = fill_capa
    rot.alignment = Alignment(horizontal="center", vertical="center")
    tot = ws["F5"]
    tot.font = Font(name=F, bold=True, size=14, color=MIDNIGHT)
    tot.fill = fill_verde
    tot.border = borda_entrada
    tot.number_format = BRL
    tot.alignment = Alignment(horizontal="right", vertical="center")

    lim = ws["E6"]
    lim.value = "Itens na cotação"
    lim.font = Font(name=F, size=9, color=CINZA)
    lim.alignment = Alignment(horizontal="center")

    ws.merge_cells("B9:F9")
    av = ws["B9"]
    av.value = (f'=IFERROR(INDEX(LISTAS!$F$1:$F${fim_cen},'
                f'MATCH($C$5&"|"&$C$6,LISTAS!$E$1:$E${fim_cen},0)),"")')
    av.font = Font(name=F, bold=True, size=10, color=MIDNIGHT)
    av.fill = fill_bege
    av.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[9].height = 32

    ws.merge_cells("B10:F10")
    nf = ws["B10"]
    nf.value = ("PREÇO POSTO FÁBRICA — o frete NÃO está incluído. Combine o frete à parte "
                "antes de fechar. Produto que não aparecer nesta lista é porque ainda não "
                "tem preço formado: peça ao Matias, não estime.")
    nf.font = Font(name=F, bold=True, size=9, color="FF8B2E2E")
    nf.fill = fill_verm
    nf.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[10].height = 30

    cabecalho(ws, 12, ["", "Produto", "Especificação", "Preço unitário", "QUANTIDADE",
                       "Total", "", "Código"])
    prod = {p.sku_key: p for p in ctx["produtos"]}
    col_cond = f"MATCH($C$7,PRECOS!$E$1:${get_column_letter(4 + len(cods))}$1,0)"
    linha = 13
    for sku in ordem_skus:
        p = prod.get(sku)
        ws.cell(row=linha, column=8, value=sku).font = Font(name=F, size=8, color=CINZA)
        ws.cell(row=linha, column=2, value=p.nome if p else "").font = Font(name=F, size=10)
        ws.cell(row=linha, column=3,
                value=(p.especificacao if p else "") or "").font = Font(name=F, size=9,
                                                                       color=CINZA)
        c = ws.cell(row=linha, column=4)
        c.value = (f'=IFERROR(INDEX(PRECOS!$E$2:${get_column_letter(4 + len(cods))}${fim},'
                   f'MATCH($H{linha}&"|"&$C$5&"|"&$C$6,PRECOS!$A$2:$A${fim},0),'
                   f'{col_cond}),"—")')
        c.number_format = BRL
        c.font = Font(name=F, bold=True, size=11, color=MIDNIGHT)
        c.alignment = Alignment(horizontal="right")
        q = ws.cell(row=linha, column=5)
        q.fill = fill_azul
        q.border = borda_entrada
        q.alignment = Alignment(horizontal="center")
        # Linha sem quantidade fica vazia, não R$ 0,00: zero afirmaria um item cotado a nada.
        tt = ws.cell(row=linha, column=6)
        tt.value = (f'=IF(OR($E{linha}="",NOT(ISNUMBER($D{linha}))),"",'
                    f'ROUND($D{linha}*$E{linha},2))')
        tt.number_format = BRL
        tt.font = Font(name=F, size=11, color=MIDNIGHT)
        for col in range(2, 7):
            ws.cell(row=linha, column=col).border = borda
        linha += 1
    fim_lista = linha - 1

    tot.value = f"=SUM(F13:F{fim_lista})"
    ws["F6"] = f'=COUNT(E13:E{fim_lista})'
    ws["F6"].font = Font(name=F, size=9, color=CINZA)
    ws["F6"].alignment = Alignment(horizontal="right")
    ws.freeze_panes = "A13"

    # ---------------------------------------------------------- COMO USAR
    cu = wb.create_sheet("COMO USAR")
    cu.sheet_view.showGridLines = False
    cu.column_dimensions["A"].width = 4
    cu.column_dimensions["B"].width = 96
    cu.merge_cells("B2:B2")
    h = cu["B2"]
    h.value = "COMO USAR"
    h.font = Font(name=F, bold=True, size=16, color=COTTON)
    h.fill = fill_capa
    h.alignment = Alignment(horizontal="center", vertical="center")
    cu.row_dimensions[2].height = 34

    blocos = [
        ("Para cotar",
         "Na aba COTAR, escolha o estado de destino, se o cliente é contribuinte de ICMS e o "
         "prazo de pagamento. A lista inteira se ajusta sozinha. Depois digite as "
         "quantidades na coluna azul QUANTIDADE — o total da cotação aparece no topo."),
        ("Contribuinte ou não contribuinte",
         "É o que mais muda o preço. Hotel, hospital e empresa que compra para uso próprio "
         "quase sempre é CONTRIBUINTE. Consumidor final e empresa sem inscrição estadual é "
         "NÃO CONTRIBUINTE, e nesse caso o preço sobe bastante em venda para fora de São "
         "Paulo — porque o imposto do estado de destino já entra no preço. Na dúvida, "
         "confirme a inscrição estadual do cliente antes de cotar. Não chute."),
        ("O frete não está aqui",
         "Todos os preços são POSTO FÁBRICA. O frete depende do peso e do tamanho do pedido "
         "inteiro, então não cabe no preço por peça. Combine o frete à parte."),
        ("Produto que não aparece",
         "Se um produto não está na lista, é porque ele ainda não tem preço formado no "
         "sistema — não porque foi esquecido. Peça ao Matias. Não use o preço de um produto "
         "parecido: as medidas e a gramatura mudam o custo."),
        ("Desconto",
         "Esta tabela é o preço recomendado. Qualquer valor ABAIXO dele precisa de aprovação, "
         "mesmo que pareça pequeno. Preço acima é livre."),
        ("Validade",
         "Os preços mudam quando o câmbio, o custo ou o imposto mudam. Confirme com o Matias "
         "antes de reusar uma tabela antiga em proposta nova."),
    ]
    linha = 4
    for titulo, texto in blocos:
        c = cu.cell(row=linha, column=2, value=titulo)
        c.font = Font(name=F, bold=True, size=12, color=COPPER)
        linha += 1
        c = cu.cell(row=linha, column=2, value=texto)
        c.font = Font(name=F, size=10, color=MIDNIGHT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        cu.row_dimensions[linha].height = 46
        linha += 2

    cu.cell(row=linha + 1, column=2,
            value=f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} · "
                  f"origem das vendas: {origem}").font = Font(name=F, size=9, italic=True,
                                                              color=CINZA)
    wb.save(SAIDA)
    return len(ordem_skus), fim


def main():
    engine = create_engine(f"sqlite:///file:{os.path.abspath(DB)}?mode=ro&uri=true")
    with Session(engine) as s:
        instalar_derivacao(s)
        grade, _pend, _int, _cen, ctx = calcular(s)
        n_skus, n_linhas = escrever(grade, ctx)
    print(f"Arquivo:        {SAIDA}")
    print(f"Produtos:       {n_skus}")
    print(f"Combinações:    {n_linhas:,}".replace(",", "."))
    print(f"Cenários:       {len(ctx['validos'])} (destino × contribuinte)")
    print("Confidencial:   nenhuma coluna de custo, margem, markup, comissão ou imposto")


if __name__ == "__main__":
    main()
