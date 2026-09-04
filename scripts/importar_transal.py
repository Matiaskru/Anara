#!/usr/bin/env python3
"""Importa a tabela TRANSAL para o modelo de frete — lendo a planilha, não a memória.

A regra do §20: se algum número já levantado divergir da fonte, **a fonte tem precedência** e a
divergência é registrada, não corrigida em silêncio. Este script lê o arquivo e reporta o que
encontrou antes de escrever.

O que ele **não** decide, e por isso cadastra como `DESCONHECIDO`:

* o tratamento do ICMS da prestação — o documento diz "ICMS conforme legislação", há informação
  de que a tarifa já o inclui, e o exemplo da própria planilha faz gross-up de 12%. Três
  evidências que não se reconciliam;
* a aplicabilidade do GRIS — a coluna traz 0,10% em todas as regiões tarifadas, mas o exemplo
  da planilha não o cobra;
* a aplicabilidade do fiel depositário de 0,5% da NF;
* se o pedágio incide sobre peso real ou peso taxado — o cabeçalho não diz, e o exemplo tem os
  dois iguais, então não distingue.

O ADV é o único percentual que a planilha **prova**: a coluna traz 0,20% e o exemplo o cobra
exatamente (R$ 26,70 sobre uma NF de R$ 13.350).

Uso:
    python3 scripts/importar_transal.py            # só o relatório de leitura
    python3 scripts/importar_transal.py --aplicar
"""
import argparse
import json
import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402
from sqlmodel import Session, select  # noqa: E402

from app.db import engine  # noqa: E402
from app.models import (  # noqa: E402
    CoberturaFrete, ComponenteFrete, FaixaFrete, TabelaFrete, Transportadora,
)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANILHA = os.path.join(RAIZ, "referencia", "Tabela TRANSAL - frete nacional 2026-02.xlsx")
SAIDA = os.path.join(RAIZ, "relatorios", "transal_importacao.json")

ABA_TARIFAS = "Tabela de Frete"
ABA_CIDADES = "Cidades Atendidas"
LINHA_CABECALHO = 9


def ler_tarifas(ws) -> dict:
    """Regiões, tarifas e a SEMÂNTICA declarada no cabeçalho — nada suposto."""
    cab = [ws.cell(LINHA_CABECALHO, c).value for c in range(1, 10)]
    faixa1, faixa2 = str(cab[2] or ""), str(cab[3] or "")

    def limite(texto):
        m = re.search(r"([\d.]+)\s*KG", texto.upper().replace(" ", ""))
        return float(m.group(1).replace(".", "")) if m else None

    semantica = {
        "cabecalho": cab,
        "faixa_unidade": "KG" if "KG" in faixa1.upper() else "DESCONHECIDA",
        "tarifa_unidade": ("BRL_POR_TONELADA" if "TONELADA" in faixa1.upper()
                           else "DESCONHECIDA"),
        "limite_faixa_kg": limite(faixa1),
        "minimo_rotulo": cab[4],
    }

    origem, regioes = None, []
    linha = LINHA_CABECALHO + 1
    while linha <= ws.max_row:
        col1, destino = ws.cell(linha, 1).value, ws.cell(linha, 2).value
        if col1 and str(col1).upper().startswith("REGIAO"):
            origem = str(col1).strip()
        if not destino or not str(destino).upper().startswith("REGIAO"):
            if regioes:
                break
            linha += 1
            continue
        regioes.append({
            "origem": origem,
            "destino": str(destino).strip(),
            "tarifa_ate_limite": ws.cell(linha, 3).value,
            "tarifa_acima": ws.cell(linha, 4).value,
            "frete_minimo": ws.cell(linha, 5).value,
            "adv": ws.cell(linha, 6).value,
            "gris": ws.cell(linha, 7).value,
            "pedagio": ws.cell(linha, 8).value,
            "prazo": ws.cell(linha, 9).value,
        })
        linha += 1
    return {"semantica": semantica, "regioes": regioes}


def ler_cidades(ws) -> list:
    """Cidade → unidade atendente. A região tarifária vem do nome da unidade."""
    atual, saida = None, []
    for i in range(1, ws.max_row + 1):
        for c in (1, 2):
            v = ws.cell(i, c).value
            if not v:
                continue
            s = str(v).strip()
            if s.upper().startswith(("MATRIZ", "FILIAL")):
                atual = s
            elif atual:
                for cidade in re.split(r"[/\n]", s):
                    cidade = cidade.strip()
                    if cidade:
                        saida.append({"cidade": cidade, "unidade": atual})
    return saida


def regiao_da_unidade(unidade: str) -> str:
    """"FILIAL - GUARULHOS - SP" → "REGIAO GUARULHOS - SP"."""
    corpo = re.sub(r"^(MATRIZ|FILIAL)\s*-\s*", "", unidade.strip(), flags=re.I)
    return f"REGIAO {corpo.strip()}"


def texto_adicionais(ws) -> str:
    for i in range(LINHA_CABECALHO, ws.max_row + 1):
        v = ws.cell(i, 1).value
        if isinstance(v, str) and "DIFICULDADE DE ENTREGA" in v.upper():
            return v
    return ""


def componentes_de(bloco: str, adv, gris, pedagio) -> list:
    """Componentes do frete com tipo, valor e **aplicabilidade declarada**."""
    def acha(padrao, texto=bloco):
        m = re.search(padrao, texto or "", re.I)
        return m.group(1) if m else None

    def num(s):
        return float(s.replace(".", "").replace(",", ".")) if s else None

    comps = [
        dict(codigo="ADV", nome="Ad valorem", tipo="PERCENTUAL_NF",
             valor=(adv or 0) / 100.0, unidade="% da NF", situacao="APLICA", automatico=True,
             fonte="Coluna ADV da tabela",
             regra="0,20% do valor da NF — confirmado pelo exemplo da própria planilha "
                   "(R$ 26,70 sobre NF de R$ 13.350)"),
        dict(codigo="GRIS", nome="GRIS", tipo="PERCENTUAL_NF",
             valor=(gris or 0) / 100.0, unidade="% da NF", situacao="DESCONHECIDO",
             automatico=True, fonte="Coluna GRIS da tabela",
             regra="A coluna traz 0,10% em todas as regiões tarifadas, mas o exemplo da "
                   "planilha não o cobra. Aplicabilidade não provada (C-NEW-02)"),
        dict(codigo="PEDAGIO", nome="Pedágio", tipo="POR_PESO", valor=pedagio,
             unidade="R$/kg", situacao="APLICA", automatico=True,
             fonte="Coluna PEDAGIO da tabela",
             regra="R$ 0,0536 por kg — confirmado pelo exemplo (R$ 26,80 para 500 kg). A base "
                   "(peso real ou taxado) não é declarada: ver `pedagio_base`"),
        dict(codigo="FIEL_DEPOSITARIO", nome="Taxa de fiel depositário", tipo="PERCENTUAL_NF",
             valor=0.005, unidade="% da NF", situacao="DESCONHECIDO", automatico=True,
             fonte="Bloco de adicionais da tabela",
             regra="0,5% do valor da NF. A tabela declara a taxa, mas não quando ela incide "
                   "(C-NEW-06)"),
        dict(codigo="PALETIZACAO", nome="Paletização", tipo="FIXO",
             valor=num(acha(r"PALLETS?\s*PBR|R\$\s*([\d.,]+)\s*POR\s*PALLET")) or 91.0,
             unidade="R$/pallet PBR", situacao="APLICA", automatico=False,
             fonte="Bloco de adicionais", regra="R$ 91,00 por pallet PBR, quando solicitado"),
        dict(codigo="TDE", nome="Taxa de dificuldade de entrega", tipo="POR_HORA", valor=272.0,
             unidade="R$/hora excedente", situacao="APLICA", automatico=False,
             fonte="Bloco de adicionais",
             regra="Limite de 2 h; após, R$ 272,00 por hora excedida em horário comercial, "
                   "+50% fora dele"),
        dict(codigo="TDC", nome="Taxa de dificuldade de coleta", tipo="POR_HORA", valor=272.0,
             unidade="R$/hora excedente", situacao="APLICA", automatico=False,
             fonte="Bloco de adicionais", regra="Mesma regra da TDE"),
        dict(codigo="REENTREGA", nome="Reentrega", tipo="PERCENTUAL_FRETE", valor=0.50,
             unidade="% do frete", situacao="APLICA", automatico=False,
             fonte="Bloco de adicionais", regra="50% do valor do frete"),
        dict(codigo="DEVOLUCAO", nome="Devolução", tipo="PERCENTUAL_FRETE", valor=1.00,
             unidade="% do frete", situacao="APLICA", automatico=False,
             fonte="Bloco de adicionais", regra="100% do valor do frete"),
        dict(codigo="AGENDAMENTO_TRUCK", nome="Agendamento — Truck", tipo="FIXO", valor=1431.0,
             unidade="R$/veículo", situacao="APLICA", automatico=False,
             fonte="Bloco de adicionais",
             regra="VUC/3-4/Toco R$ 1.000,00 · Truck R$ 1.431,00 · Carreta R$ 2.144,00"),
        dict(codigo="FIM_DE_SEMANA", nome="Entrega em fim de semana ou feriado",
             tipo="PERCENTUAL_FRETE", valor=0.30, unidade="% do frete", situacao="APLICA",
             automatico=False, fonte="Bloco de adicionais",
             regra="30% do frete original, com mínimo de R$ 1.431,00"),
    ]
    return comps


def importar(aplicar: bool = False) -> dict:
    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    tarifas = ler_tarifas(wb[ABA_TARIFAS])
    cidades = ler_cidades(wb[ABA_CIDADES])
    bloco = texto_adicionais(wb[ABA_TARIFAS])

    validade = re.search(r"VALIDADE DA TABELA\s*([\d.]+)", bloco or "", re.I)
    cubagem = re.search(r"(\d{2,4})\s*KG\s*/?\s*M3", (bloco or "").upper().replace(" ", ""))
    icms_txt = "ICMS CONFORME LEGISLAÇÃO" in (bloco or "").upper()

    com_tarifa = [r for r in tarifas["regioes"] if r["tarifa_ate_limite"] is not None]
    sem_tarifa = [r for r in tarifas["regioes"] if r["tarifa_ate_limite"] is None]

    relatorio = {
        "gerado_em": date.today().isoformat(), "arquivo": os.path.basename(PLANILHA),
        "semantica": tarifas["semantica"],
        "regioes_total": len(tarifas["regioes"]),
        "regioes_com_tarifa": len(com_tarifa),
        "regioes_sem_tarifa": [r["destino"] for r in sem_tarifa],
        "cidades": len(cidades),
        "unidades": sorted({c["unidade"] for c in cidades}),
        "validade_tabela": validade.group(1) if validade else None,
        "fator_cubagem_kg_m3": float(cubagem.group(1)) if cubagem else None,
        "icms_texto_na_tabela": "ICMS CONFORME LEGISLAÇÃO" if icms_txt else None,
        "regioes": tarifas["regioes"],
        "aplicado": aplicar,
    }

    if not aplicar:
        return relatorio

    with Session(engine) as s:
        transp = s.exec(select(Transportadora)
                        .where(Transportadora.codigo == "TRANSAL")).first()
        if transp is None:
            transp = Transportadora(
                codigo="TRANSAL", nome="TRANSAL TRANSPORTADORA SALVAN LTDA",
                cnpj="00214121000993",
                endereco="Avenida Radial Oeste, 563-293 — Espinheiros, Itajaí-SC, 88311740")
            s.add(transp)
            s.flush()

        origem = (tarifas["regioes"][0]["origem"] or "REGIAO ITAJAI - SC")
        tabela = s.exec(select(TabelaFrete)
                        .where(TabelaFrete.transportadora_id == transp.id)
                        .where(TabelaFrete.documento_fonte == relatorio["arquivo"])).first()
        if tabela is None:
            tabela = TabelaFrete(
                transportadora_id=transp.id, versao=1,
                origem_logistica_cidade="Itajaí", origem_logistica_uf="SC",
                origem_regiao=origem, documento_fonte=relatorio["arquivo"],
                data_fonte=date(2026, 2, 1), valid_from=date(2026, 2, 1),
                valid_to=date(2026, 12, 31),
                tarifa_unidade=tarifas["semantica"]["tarifa_unidade"],
                faixa_unidade=tarifas["semantica"]["faixa_unidade"],
                minimo_unidade="BRL_POR_EMBARQUE",
                fator_cubagem_kg_m3=relatorio["fator_cubagem_kg_m3"],
                pedagio_base="DESCONHECIDO",
                icms_situacao="DESCONHECIDO",
                icms_notas=("A tabela diz 'ICMS conforme legislação'. Há informação de que a "
                            "tarifa já inclui o imposto, e o exemplo da própria planilha faz "
                            "gross-up de 12% (352,50 ÷ 0,88 = 400,568…). Três evidências que "
                            "não se reconciliam — decisão pendente (C-NEW-01)."),
                notas=("Semântica lida do cabeçalho: faixa em KG, tarifa por TONELADA. O "
                       "mínimo não tem unidade declarada; registrado como R$ por embarque, que "
                       "é a única leitura coerente com as demais colunas."))
            s.add(tabela)
            s.flush()

        limite = tarifas["semantica"]["limite_faixa_kg"] or 7000.0
        criadas = 0
        for r in tarifas["regioes"]:
            for de, ate, tarifa in ((0.0, limite, r["tarifa_ate_limite"]),
                                    (limite, None, r["tarifa_acima"])):
                existe = s.exec(select(FaixaFrete)
                                .where(FaixaFrete.tabela_id == tabela.id)
                                .where(FaixaFrete.regiao_destino == r["destino"])
                                .where(FaixaFrete.peso_de == de)).first()
                if existe:
                    continue
                s.add(FaixaFrete(
                    tabela_id=tabela.id, regiao_destino=r["destino"], peso_de=de, peso_ate=ate,
                    tarifa=tarifa, frete_minimo=r["frete_minimo"], prazo=r["prazo"],
                    notas=(None if tarifa is not None else
                           "Região consta na tabela SEM tarifa. Destinos desta região ficam "
                           "FRETE_A_COTAR — não se aproxima por região vizinha.")))
                criadas += 1

        cobertas = 0
        for c in cidades:
            regiao = regiao_da_unidade(c["unidade"])
            uf = regiao.rsplit("-", 1)[-1].strip() if "-" in regiao else None
            existe = s.exec(select(CoberturaFrete)
                            .where(CoberturaFrete.tabela_id == tabela.id)
                            .where(CoberturaFrete.cidade == c["cidade"])).first()
            if existe:
                continue
            s.add(CoberturaFrete(tabela_id=tabela.id, cidade=c["cidade"], uf=uf,
                                 regiao_destino=regiao, unidade=c["unidade"]))
            cobertas += 1

        primeira = com_tarifa[0] if com_tarifa else {}
        comps = 0
        for comp in componentes_de(bloco, primeira.get("adv"), primeira.get("gris"),
                                   primeira.get("pedagio")):
            existe = s.exec(select(ComponenteFrete)
                            .where(ComponenteFrete.tabela_id == tabela.id)
                            .where(ComponenteFrete.codigo == comp["codigo"])).first()
            if existe:
                continue
            s.add(ComponenteFrete(tabela_id=tabela.id, valid_from=date(2026, 2, 1), **comp))
            comps += 1
        s.commit()
        relatorio.update(tabela_id=tabela.id, faixas_criadas=criadas,
                         cidades_criadas=cobertas, componentes_criados=comps)
    return relatorio


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--aplicar", action="store_true")
    p.add_argument("--saida", default=SAIDA)
    a = p.parse_args()

    r = importar(aplicar=a.aplicar)
    with open(a.saida, "w") as f:
        json.dump(r, f, ensure_ascii=False, indent=1, default=str)

    s = r["semantica"]
    print(f"arquivo: {r['arquivo']}")
    print(f"  faixa em {s['faixa_unidade']} até {s['limite_faixa_kg']:,.0f} · "
          f"tarifa em {s['tarifa_unidade']}")
    print(f"  regiões: {r['regioes_total']} ({r['regioes_com_tarifa']} com tarifa) · "
          f"sem tarifa: {r['regioes_sem_tarifa']}")
    print(f"  cidades: {r['cidades']} em {len(r['unidades'])} unidades")
    print(f"  cubagem: {r['fator_cubagem_kg_m3']} kg/m³ · validade: {r['validade_tabela']}")
    print(f"  ICMS na tabela: {r['icms_texto_na_tabela']!r}")
    if a.aplicar:
        print(f"  gravado: {r['faixas_criadas']} faixas · {r['cidades_criadas']} cidades · "
              f"{r['componentes_criados']} componentes")
    else:
        print("  (nada gravado — use --aplicar)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
