#!/usr/bin/env python3
"""Lista de preços completa em Excel — SKU × destino × contribuinte × condição de pagamento.

A planilha **não faz conta financeira**. Todo preço sai do motor determinístico
(`pricing_service` + `pricing_engine`), já quantizado, e é gravado como número. O Excel só
faz `MATCH`/`INDEX` numa grade pronta. É o que impede a régua da planilha divergir da régua
do sistema — o erro clássico de reimplementar ICMS e encargo em fórmula de célula.

Abas:

* **CONSULTA** — três listas (destino, contribuinte, pagamento) e a lista de preços inteira
  para o cenário escolhido. É a aba que vai para o cliente: não tem custo, margem, markup,
  comissão, imposto nem fornecedor.
* **GRADE** — a matriz pré-calculada que a CONSULTA consulta. Uma linha por
  SKU × estado × contribuinte; uma coluna por condição de pagamento.
* **PENDENCIAS** — o que **não** formou preço, com o motivo. Não é sobra: é a lista de
  perguntas a fazer, e por isso vem junto.
* **INTERNO** — margem, custo NET e comissão do cenário-base, para conferência.
  **Não enviar ao cliente.**
* **PREMISSAS** — câmbio, encargos, origem e data. Responde "de onde veio este número".

Cenário sem regra fiscal resolvida **não vira linha**. Não existe preço "aproximado" aqui:
combinação bloqueada aparece em PENDENCIAS com o motivo que o motor deu.

Rodar:  python3 scripts/gerar_lista_precos.py
"""
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime

sys.path.insert(0, os.path.expanduser("~/Anara-Cotacao"))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import create_engine
from sqlmodel import Session, select

import app.config_service as cfg
import app.pricing_service as ps
from app.dinheiro import para_float
from app.models import (AliquotaInterestadual, CoberturaFrete, ComponenteFrete,
                        CondicaoPagamento, Cotacao, EstadoFiscal, FaixaFrete, Fornecedor,
                        Produto, RegraFcp, RegraFiscalVenda, TabelaFrete)
from app.pricing_engine import calcular_por_margem

DB = os.path.expanduser("~/Anara-Cotacao/data/anara.db")
SAIDA = os.path.expanduser("~/Downloads/Lista de Precos Anara.xlsx")

# Identidade visual da marca, igual à da calculadora e do PDF.
MIDNIGHT, COPPER, SAND = "FF30354F", "FFC49281", "FFD6C8BE"
COTTON, CINZA = "FFF9F6F2", "FF6B6E7C"
AZUL, VERDE, BEGE, VERMELHO = "FFD8E4F5", "FFE3EFE7", "FFF4EFEA", "FFF6E0DC"
F = "Calibri"
BRL = '"R$ "#,##0.00'
PCT = "0.0%"

fill_capa = PatternFill("solid", fgColor=MIDNIGHT)
fill_cab = PatternFill("solid", fgColor=MIDNIGHT)
fill_secao = PatternFill("solid", fgColor=COPPER)
fill_azul = PatternFill("solid", fgColor=AZUL)
fill_bege = PatternFill("solid", fgColor=BEGE)
fill_verm = PatternFill("solid", fgColor=VERMELHO)
lado = Side(style="thin", color=SAND)
borda = Border(left=lado, right=lado, top=lado, bottom=lado)
grosso = Side(style="medium", color=COPPER)
borda_entrada = Border(left=grosso, right=grosso, top=grosso, bottom=grosso)


def cabecalho(ws, linha, titulos, larguras=None):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(name=F, bold=True, color=COTTON, size=10)
        c.fill = fill_cab
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
    ws.row_dimensions[linha].height = 30
    for i, w in enumerate(larguras or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ===========================================================================
# 1. Premissas DERIVADAS — a camada que o motor se recusa a assumir sozinho
# ===========================================================================
# O motor bloqueia a venda a não contribuinte fora de SP e RJ por duas premissas que ninguém
# confirmou: a base interna do destino (`icms_interno_base`) e a incidência de FCP. As duas
# são deliberadas — `fiscal_rules` diz, com todas as letras, que `EstadoFiscal.fem` "não
# alimenta o motor", e que ausência de regra de FCP não é 0%.
#
# Aqui elas são DERIVADAS, a pedido, para que a lista saia completa. A derivação é auditável e
# foi conferida contra os dados antes de ser usada:
#
#   base_simples = interna − 4%                    → confere em 27/27 estados
#   base_dupla   = (interna − 4%) / (1 − interna)  → confere em 27/27 estados
#   icms_interno_base = interna − fem              → confere no RJ, único estado já
#                                                    determinado à mão (22% − 2% = 20%)
#
# As três identidades só fecham se `aliquota_interna` for a alíquota CHEIA, com o FCP dentro, e
# `fem` for a parcela de FCP. Quatro estados têm fem (BA, PE, PI, RJ, todos 2%); nos outros 23
# a base é a própria alíquota interna.
#
# ATENÇÃO ao limite desta validação. Ela prova a DECOMPOSIÇÃO aritmética — que a alíquota
# interna se separa em base + FCP — e nada além disso. Não prova a INCIDÊNCIA, que é outra
# pergunta: o FCP não incide sobre tudo que entra num estado, depende do produto, e é por isso
# que `RegraFcp` existe como linha cadastrada em vez de coluna por UF. A conferência de que a
# carga total do não-contribuinte bate com a alíquota interna do destino é **circular**: com
# base = interna − fem, `interestadual + (base − interestadual) + fem` devolve `interna` por
# construção, em qualquer estado, mesmo que o FCP não incidisse. E o RJ também não decide, por
# ser justamente o único estado que JÁ tinha `RegraFcp` cadastrada e base determinada — a
# derivação não é usada lá.
#
# Ou seja: para os 26 estados derivados, esta planilha **assume** que o FCP se aplica aos
# produtos da Anara na alíquota de `fem` (e que não se aplica onde `fem = 0`). Isso precisa
# de confirmação do contador, e é uma pergunta diferente de "qual é a base interna".
#
# Isto **não é gravado no banco**. Cadastrar premissa é ato versionado, com preview, trilha e
# autorização própria — não efeito colateral de gerar planilha. Aqui a derivação vive só na
# memória deste processo e sai carimbada em toda aba que a usa.
DERIVACAO = ("icms_interno_base = aliquota_interna − fem · "
             "FCP assumido = fem (incidência sobre o produto NÃO verificada)")


def derivar_premissas(estados):
    """Preenche a base interna dos destinos que o motor deixaria bloqueados.

    Devolve a lista do que foi derivado, para a planilha poder mostrar estado por estado o
    que é dado cadastrado e o que é inferência.
    """
    derivados = []
    for e in estados:
        cadastrada = getattr(e, "icms_interno_base", None)
        fem = float(getattr(e, "fem", 0.0) or 0.0)
        interna = float(e.aliquota_interna or 0.0)
        if cadastrada is not None:
            derivados.append({"uf": e.uf, "interna": interna, "fem": fem,
                              "base": float(cadastrada), "fonte": "CADASTRADO"})
            continue
        e.icms_interno_base = interna - fem
        derivados.append({"uf": e.uf, "interna": interna, "fem": fem,
                          "base": interna - fem, "fonte": "DERIVADO"})
    return derivados


def regras_fcp_derivadas(estados, reais):
    """Uma linha de FCP por UF, a partir da coluna `fem`.

    UF que já tem regra cadastrada é preservada — a cadastrada vence, e a prioridade mais
    baixa garante isso mesmo se as duas alcançarem o item.
    """
    ja_tem = {(r.uf_destino or "").upper() for r in reais}
    novas = list(reais)
    for e in estados:
        if e.uf.upper() in ja_tem:
            continue
        fem = float(getattr(e, "fem", 0.0) or 0.0)
        novas.append(RegraFcp(
            id=900000 + (e.id or 0), uf_destino=e.uf, fcp_pct=fem, prioridade=900,
            situacao="APLICA" if fem > 0 else "NAO_APLICA",
            regra=f"DERIVADO da coluna EstadoFiscal.fem ({fem:.2%})",
            fonte=f"Derivação da lista de preços — {DERIVACAO}",
            valid_from=None, ativo=True))
    return novas


def instalar_derivacao(session):
    """Faz `pricing_service` enxergar as premissas derivadas, sem tocar no banco.

    Substitui só a **fiação** de `fiscal_do_item` — que estados e que regras de FCP entram.
    A matemática continua onde deve estar: em `resolver_fiscal_item`, intocada. Nada é
    gravado; a sessão é somente leitura e o autoflush fica desligado por garantia.
    """
    from app.fiscal_rules import normalizar_uf, resolver_fiscal_item

    session.autoflush = False
    estados = session.exec(select(EstadoFiscal)).all()
    derivados = derivar_premissas(estados)
    fcp = regras_fcp_derivadas(estados, session.exec(select(RegraFcp)).all())
    explicitas = session.exec(select(RegraFiscalVenda)).all()
    aliquotas = session.exec(select(AliquotaInterestadual)).all()

    def fiscal_do_item(sessao, cotacao, produto=None):
        uf_origem, fonte_origem = ps.uf_origem_fiscal(sessao, cotacao, produto)
        uf_destino = normalizar_uf(estados, getattr(cotacao, "estado_destino", None))
        origem_fiscal, fonte_natureza = ps.origem_fiscal_do_produto(sessao, produto)
        finalidade, fonte_finalidade = ps.finalidade_da_operacao(sessao, cotacao)
        r = resolver_fiscal_item(
            explicitas, estados, aliquotas,
            uf_origem=uf_origem, uf_destino=uf_destino, origem_fiscal=origem_fiscal,
            contribuinte=getattr(cotacao, "contribuinte_icms", None),
            finalidade=finalidade, ncm=getattr(produto, "ncm", None),
            produto_id=getattr(produto, "id", None),
            familia=getattr(produto, "familia", None), regras_fcp=fcp)
        r.avisos.append(f"origem fiscal: {fonte_origem}")
        r.avisos.append(f"natureza da mercadoria: {fonte_natureza}")
        r.avisos.append(f"finalidade: {fonte_finalidade}")
        return r

    ps.fiscal_do_item = fiscal_do_item
    return derivados


def chave_regiao(texto):
    """Mesma normalização que `frete_service._chave` usa para casar região com cobertura.

    A tabela da TRANSAL grafa a mesma região de dois jeitos — "REGIAO ITAJAÍ - SC" na
    cobertura, "REGIAO ITAJAI - SC" nas faixas. O motor normaliza e casa; a planilha
    precisa fazer o mesmo, senão mostra dois nomes para a mesma região.
    """
    import unicodedata
    base = "".join(c for c in unicodedata.normalize("NFD", texto or "")
                   if unicodedata.category(c) != "Mn")
    return base.strip().upper()


def frete_da_base(session):
    """A tabela TRANSAL como ela está — tarifas, faixas, cobertura e adicionais.

    O frete **não vira coluna de preço unitário**, e isso não é premissa faltando: a tarifa é
    por tonelada, com mínimo por embarque e adicionais sobre o valor da NF. O frete por peça
    de um pedido de 10 peças e o de um de 10.000 diferem em ordem de grandeza. Só existe frete
    de um EMBARQUE — por isso a tabela vem inteira, para consulta, e o rateio por item
    acontece na cotação, onde o peso e o valor do pedido são conhecidos.
    """
    tabelas = session.exec(select(TabelaFrete)).all()
    return {
        "tabelas": tabelas,
        "faixas": session.exec(select(FaixaFrete)).all(),
        "componentes": session.exec(select(ComponenteFrete)).all(),
        "cobertura": session.exec(select(CoberturaFrete)).all(),
    }


# ===========================================================================
# 2. Motor — a grade é calculada em Python, nunca no Excel
# ===========================================================================
def calcular(session):
    """Devolve (grade, pendencias, interno, cenarios, contexto).

    `grade[(sku, estado, contrib)][codigo_condicao] = preco`. Só entram combinações que o
    motor resolveu: bloqueio vira pendência, não vira número.
    """
    estados = sorted(session.exec(select(EstadoFiscal)).all(), key=lambda e: e.estado)
    fornecedores = {f.id: f.nome for f in session.exec(select(Fornecedor)).all()}

    # Só condições com encargo CONFIRMADO. `SINAL30+30/60/90` e `CARTAO` estão cadastradas
    # sem encargo confirmado — entrariam como 0%, que é exatamente o fallback silencioso
    # que o projeto proíbe.
    condicoes = [c for c in sorted(session.exec(select(CondicaoPagamento)).all(),
                                   key=lambda c: c.ordem)
                 if c.ativo and c.encargo_confirmado]

    produtos = sorted(session.exec(select(Produto).where(Produto.ativo == True)).all(),  # noqa: E712
                      key=lambda p: (p.familia or "zzz", p.nome or ""))
    finalidade = cfg.txt(session, "fiscal_finalidade_padrao", "USO_CONSUMO")
    origem_nome = cfg.txt(session, "catalogo_origem", "São Paulo")
    origem_uf = next((e.uf for e in estados if e.estado == origem_nome), "SP")

    grade, pendencias, interno = {}, [], []
    cenarios = {}
    motivos = Counter()

    # --- que pares (estado, contribuinte) o fiscal resolve? -----------------------------
    # Resolvido uma vez com um produto-sonda: a regra fiscal depende do par de UF, da
    # natureza e do contribuinte — não do SKU. Perguntar 295 vezes daria a mesma resposta.
    sonda = next((p for p in produtos if p.custo_unitario), None)
    if sonda is None:
        raise SystemExit("Nenhum produto ativo com custo. Nada a exportar.")

    fiscal_detalhe = []
    por_natureza = {}
    for p in produtos:
        nat, _ = ps.origem_fiscal_do_produto(session, p)
        if nat and nat not in por_natureza:
            por_natureza[nat] = p

    for e in estados:
        for contrib in (True, False):
            virtual = Cotacao(cliente_id=0, estado_origem=origem_nome,
                              uf_origem_fiscal=origem_uf, estado_destino=e.estado,
                              contribuinte_icms=contrib, finalidade=finalidade,
                              condicao_pagamento=condicoes[0].codigo)
            regras, ctx = ps.regras_da_cotacao(session, virtual, sonda)
            motivo = None if regras is not None else (ctx.get("motivo_bloqueio") or "sem regra")
            cenarios[(e.estado, contrib)] = motivo
            if motivo:
                motivos[motivo.split(":")[0][:80]] += 1

            # A alíquota depende da NATUREZA da mercadoria: o mesmo par de UF dá 4% para a
            # KTC (importada) e 12% para Daune/Decor (nacional). Uma linha por natureza.
            for nat, amostra in sorted(por_natureza.items()):
                f = ps.fiscal_do_item(session, virtual, amostra)
                fiscal_detalhe.append({
                    "uf": e.uf, "estado": e.estado, "contribuinte": contrib,
                    "natureza": nat,
                    "icms_total": para_float(f.icms_pct),
                    "interestadual": para_float(getattr(f, "aliquota_interestadual", None)),
                    "difal": para_float(f.difal_pct),
                    "fcp": para_float(f.fcp_pct),
                    "quem_recolhe": f.difal_responsavel or "—",
                    "na_margem": "SIM" if getattr(f, "difal_entra_na_margem", False) else "NÃO",
                    "interna_destino": float(e.aliquota_interna or 0),
                    "base_interna": float(getattr(e, "icms_interno_base", 0) or 0),
                    "fem": float(getattr(e, "fem", 0) or 0),
                    "status": f.status, "regra": (f.regra or "")[:260],
                })

    validos = [(est, c) for (est, c), motivo in cenarios.items() if motivo is None]

    # --- preço por SKU × cenário válido × condição --------------------------------------
    for p in produtos:
        if not p.custo_unitario:
            pendencias.append({
                "sku": p.sku_key, "nome": p.nome, "familia": p.familia,
                "fornecedor": fornecedores.get(p.fornecedor_id, "—"),
                "tipo": "SEM CUSTO",
                "detalhe": "Nenhum custo cadastrado: o produto não forma preço em cenário nenhum.",
            })
            continue

        margem = ps.margem_padrao(session, p)
        for est, contrib in validos:
            virtual_base = None
            for cond in condicoes:
                virtual = Cotacao(cliente_id=0, estado_origem=origem_nome,
                                  uf_origem_fiscal=origem_uf, estado_destino=est,
                                  contribuinte_icms=contrib, finalidade=finalidade,
                                  condicao_pagamento=cond.codigo)
                regras, ctx = ps.regras_da_cotacao(session, virtual, p)
                if regras is None:
                    # O cenário resolveu para a sonda mas não para este SKU — quase sempre
                    # natureza da mercadoria. Vira pendência do SKU, não preço chutado.
                    pendencias.append({
                        "sku": p.sku_key, "nome": p.nome, "familia": p.familia,
                        "fornecedor": fornecedores.get(p.fornecedor_id, "—"),
                        "tipo": f"BLOQUEADO — {est} / {'contribuinte' if contrib else 'não contribuinte'}",
                        "detalhe": (ctx.get("motivo_bloqueio") or "sem regra")[:300],
                    })
                    continue
                r = calcular_por_margem(p.custo_unitario, 1.0, margem.margem_pct,
                                        regras, preco_base=p.preco_base)
                grade.setdefault((p.sku_key, est, contrib), {})[cond.codigo] = \
                    para_float(r.preco_negociado)
                if virtual_base is None:
                    virtual_base = r

            # Conferência interna: um cenário por SKU basta para validar margem/comissão.
            if virtual_base is not None and (est, contrib) == validos[0]:
                custo = ps.custo_net(session, p)
                interno.append({
                    "sku": p.sku_key, "nome": p.nome, "familia": p.familia,
                    "fornecedor": fornecedores.get(p.fornecedor_id, "—"),
                    "custo_net": para_float(custo.get("net_brl")),
                    "margem_alvo": para_float(margem.margem_pct),
                    "margem_regra": margem.regra,
                    "preco": para_float(virtual_base.preco_negociado),
                    "margem_real": para_float(virtual_base.margem_liquida),
                    "comissao": para_float(virtual_base.comissao),
                    "impostos": para_float(virtual_base.impostos),
                    "lucro": para_float(virtual_base.lucro),
                    "confianca": p.custo_confianca or "—",
                    "revisao": "SIM" if p.precisa_revisao else "",
                })

    # Cenários bloqueados viram pendência própria — uma linha por cenário, não por SKU,
    # senão 25 estados × 295 SKUs afogariam a aba.
    for (est, contrib), motivo in sorted(cenarios.items()):
        if motivo:
            pendencias.append({
                "sku": "—", "nome": f"TODO O CATÁLOGO — {est}", "familia": "—",
                "fornecedor": "—",
                "tipo": f"CENÁRIO BLOQUEADO — {'contribuinte' if contrib else 'não contribuinte'}",
                "detalhe": motivo[:300],
            })

    contexto = {
        "origem_nome": origem_nome, "origem_uf": origem_uf, "finalidade": finalidade,
        "estados": estados, "condicoes": condicoes, "produtos": produtos,
        "validos": validos, "cenarios": cenarios, "motivos": motivos,
        "premissas": ps.pinar_premissas(session), "fornecedores": fornecedores,
        "fiscal_detalhe": fiscal_detalhe, "frete": frete_da_base(session),
    }
    return grade, pendencias, interno, cenarios, contexto


# ===========================================================================
# 2. Escrita da planilha
# ===========================================================================
def escrever(grade, pendencias, interno, cenarios, ctx):
    wb = Workbook()
    condicoes = ctx["condicoes"]
    cods = [c.codigo for c in condicoes]

    # ---------------------------------------------------------------- GRADE
    gr = wb.active
    gr.title = "GRADE"
    cabecalho(gr, 1, ["CHAVE", "SKU", "Destino", "Contribuinte"] + cods,
              [34, 20, 22, 13] + [17] * len(cods))
    linha = 2
    skus_com_preco = []
    vistos = set()
    for (sku, est, contrib), precos in sorted(grade.items()):
        chave = f"{sku}|{est}|{'SIM' if contrib else 'NAO'}"
        gr.cell(row=linha, column=1, value=chave)
        gr.cell(row=linha, column=2, value=sku)
        gr.cell(row=linha, column=3, value=est)
        gr.cell(row=linha, column=4, value="SIM" if contrib else "NAO")
        for i, cod in enumerate(cods, start=5):
            c = gr.cell(row=linha, column=i, value=precos.get(cod))
            c.number_format = BRL
        if sku not in vistos:
            vistos.add(sku)
            skus_com_preco.append(sku)
        linha += 1
    gr.freeze_panes = "B2"
    fim_grade = linha - 1

    # ------------------------------------------------------------- CONSULTA
    ws = wb.create_sheet("CONSULTA", 0)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:F2")
    t = ws["B2"]
    t.value = "ANARA — LISTA DE PREÇOS"
    t.font = Font(name=F, bold=True, size=18, color=COTTON)
    t.fill = fill_capa
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 40

    ws.merge_cells("B3:F3")
    s = ws["B3"]
    s.value = ("Escolha destino, tipo de cliente e condição de pagamento nas três caixas "
               "azuis. A lista inteira muda sozinha. O que não tem preço está em PENDENCIAS, "
               "com o motivo.")
    s.font = Font(name=F, size=10, italic=True, color=CINZA)
    s.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[3].height = 26

    entradas = [("B5", "C5", "Estado de destino", "São Paulo"),
                ("B6", "C6", "Cliente é contribuinte de ICMS?", "SIM"),
                ("B7", "C7", "Condição de pagamento", cods[0])]
    for rot, cel, texto, padrao in entradas:
        r = ws[rot]
        r.value = texto
        r.font = Font(name=F, bold=True, size=11, color=MIDNIGHT)
        r.alignment = Alignment(horizontal="right", vertical="center")
        c = ws[cel]
        c.value = padrao
        c.font = Font(name=F, bold=True, size=12, color=MIDNIGHT)
        c.fill = fill_azul
        c.border = borda_entrada
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[int(rot[1:])].height = 24

    # Listas de apoio, numa aba escondida.
    lst = wb.create_sheet("LISTAS")
    for i, e in enumerate(ctx["estados"], start=1):
        lst.cell(row=i, column=1, value=e.estado)
    lst.cell(row=1, column=2, value="SIM")
    lst.cell(row=2, column=2, value="NAO")
    for i, cod in enumerate(cods, start=1):
        lst.cell(row=i, column=3, value=cod)
    # Mapa de cenário → motivo do bloqueio, para a CONSULTA avisar em vez de mentir.
    for i, ((est, contrib), motivo) in enumerate(sorted(cenarios.items()), start=1):
        lst.cell(row=i, column=5, value=f"{est}|{'SIM' if contrib else 'NAO'}")
        lst.cell(row=i, column=6, value=motivo or "OK")
    fim_cenarios = len(cenarios)
    lst.sheet_state = "hidden"

    n_est = len(ctx["estados"])
    for cel, formula in (("C5", f"LISTAS!$A$1:$A${n_est}"),
                         ("C6", "LISTAS!$B$1:$B$2"),
                         ("C7", f"LISTAS!$C$1:$C${len(cods)}")):
        dv = DataValidation(type="list", formula1=f"={formula}", allow_blank=False,
                            showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(ws[cel])

    # Aviso do cenário: lê o motivo real do motor, não inventa texto.
    ws.merge_cells("B9:F9")
    av = ws["B9"]
    av.value = ('=IFERROR(IF(INDEX(LISTAS!$F$1:$F$' + str(fim_cenarios) + ',MATCH($C$5&"|"&$C$6,'
                'LISTAS!$E$1:$E$' + str(fim_cenarios) + ',0))="OK",'
                '"Cenário liberado — preços abaixo.",'
                '"CENÁRIO BLOQUEADO: "&INDEX(LISTAS!$F$1:$F$' + str(fim_cenarios) +
                ',MATCH($C$5&"|"&$C$6,LISTAS!$E$1:$E$' + str(fim_cenarios) + ',0))),'
                '"Cenário não cadastrado.")')
    av.font = Font(name=F, bold=True, size=10, color=MIDNIGHT)
    av.fill = fill_bege
    av.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[9].height = 34

    # O que o preço NÃO contém. Fica na primeira tela, não numa nota de rodapé: é a diferença
    # entre o número aqui e o que o cliente vai efetivamente pagar.
    ws.merge_cells("B10:F10")
    nf = ws["B10"]
    nf.value = ("PREÇO EX-FRETE. O frete não está incluído — calcule na aba FRETE e some por "
                "fora. Em venda a CONTRIBUINTE de outro estado, o DIFAL é recolhido pelo "
                "cliente e também não está aqui (veja a aba FISCAL).")
    nf.font = Font(name=F, bold=True, size=9, color="FF8B2E2E")
    nf.fill = fill_verm
    nf.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[10].height = 30

    cabecalho(ws, 11, ["", "Produto", "Especificação", "Preço unitário", "Fornecedor",
                       "", "Código"])
    # A coluna do código é a âncora do MATCH: precisa existir, não precisa ser vista. O
    # `sku_key` repete nome e especificação, então exibi-lo só suja a lista que vai ao
    # cliente. Fica em G, longe dos rótulos das caixas de seleção — que moram em B — e
    # oculta; fórmula que aponta para coluna oculta continua calculando.
    for col, w in zip("ABCDEF", [4, 44, 44, 18, 22, 3]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["G"].hidden = True

    prod_por_sku = {p.sku_key: p for p in ctx["produtos"]}
    col_cond = f"MATCH($C$7,GRADE!$E$1:${get_column_letter(4 + len(cods))}$1,0)"
    linha = 12
    for sku in skus_com_preco:
        p = prod_por_sku.get(sku)
        ws.cell(row=linha, column=7, value=sku).font = Font(name=F, size=8, color=CINZA)
        ws.cell(row=linha, column=2, value=p.nome if p else "").font = Font(name=F, size=10)
        ws.cell(row=linha, column=3,
                value=(p.especificacao if p else "") or "").font = Font(name=F, size=9,
                                                                       color=CINZA)
        c = ws.cell(row=linha, column=4)
        c.value = (f'=IFERROR(INDEX(GRADE!$E$2:${get_column_letter(4 + len(cods))}${fim_grade},'
                   f'MATCH($G{linha}&"|"&$C$5&"|"&$C$6,GRADE!$A$2:$A${fim_grade},0),'
                   f'{col_cond}),"—")')
        c.number_format = BRL
        c.font = Font(name=F, bold=True, size=11, color=MIDNIGHT)
        c.alignment = Alignment(horizontal="right")
        ws.cell(row=linha, column=5,
                value=ctx["fornecedores"].get(p.fornecedor_id, "—") if p else "—").font = \
            Font(name=F, size=9, color=CINZA)
        for col in range(2, 6):
            ws.cell(row=linha, column=col).border = borda
        linha += 1
    ws.freeze_panes = "A12"

    # ---------------------------------------------------------- PENDENCIAS
    pd = wb.create_sheet("PENDENCIAS")
    pd.sheet_view.showGridLines = False
    pd.merge_cells("A1:E1")
    h = pd["A1"]
    h.value = ("O QUE NÃO FORMA PREÇO — e por quê. Cada linha aqui é uma informação que "
               "falta, não um erro do sistema.")
    h.font = Font(name=F, bold=True, size=11, color=COTTON)
    h.fill = fill_capa
    h.alignment = Alignment(horizontal="left", vertical="center")
    pd.row_dimensions[1].height = 30
    cabecalho(pd, 2, ["Tipo", "SKU", "Produto", "Fornecedor", "Motivo"],
              [42, 20, 40, 20, 90])
    linha = 3
    for p in sorted(pendencias, key=lambda x: (x["tipo"], x["sku"])):
        pd.cell(row=linha, column=1, value=p["tipo"]).fill = fill_verm
        pd.cell(row=linha, column=2, value=p["sku"])
        pd.cell(row=linha, column=3, value=p["nome"])
        pd.cell(row=linha, column=4, value=p["fornecedor"])
        cel = pd.cell(row=linha, column=5, value=p["detalhe"])
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        linha += 1
    pd.freeze_panes = "A3"
    pd.auto_filter.ref = f"A2:E{linha - 1}"

    # ------------------------------------------------------------- INTERNO
    it = wb.create_sheet("INTERNO")
    it.sheet_view.showGridLines = False
    it.merge_cells("A1:L1")
    h = it["A1"]
    h.value = "USO INTERNO — custo, margem e comissão. NÃO ENVIAR AO CLIENTE."
    h.font = Font(name=F, bold=True, size=12, color=COTTON)
    h.fill = PatternFill("solid", fgColor="FF8B2E2E")
    h.alignment = Alignment(horizontal="center", vertical="center")
    it.row_dimensions[1].height = 28
    est0, contrib0 = ctx["validos"][0]
    it.merge_cells("A2:L2")
    it["A2"] = (f"Cenário de conferência: {ctx['origem_nome']} → {est0} · "
                f"{'contribuinte' if contrib0 else 'não contribuinte'} · "
                f"{cods[0]} · finalidade {ctx['finalidade']}")
    it["A2"].font = Font(name=F, italic=True, size=9, color=CINZA)
    cabecalho(it, 3, ["SKU", "Produto", "Família", "Fornecedor", "Custo NET",
                      "Margem alvo", "Preço", "Margem real", "Comissão", "Impostos",
                      "Lucro", "Confiança do custo", "Precisa revisão"],
              [20, 38, 20, 18, 14, 12, 14, 12, 13, 13, 13, 18, 14])
    linha = 4
    for r in interno:
        vals = [r["sku"], r["nome"], r["familia"], r["fornecedor"], r["custo_net"],
                r["margem_alvo"], r["preco"], r["margem_real"], r["comissao"],
                r["impostos"], r["lucro"], r["confianca"], r["revisao"]]
        for i, v in enumerate(vals, start=1):
            c = it.cell(row=linha, column=i, value=v)
            if i in (5, 7, 9, 10, 11):
                c.number_format = BRL
            if i in (6, 8):
                c.number_format = PCT
        if r["revisao"]:
            it.cell(row=linha, column=13).fill = fill_verm
        linha += 1
    it.freeze_panes = "A4"
    it.auto_filter.ref = f"A3:M{linha - 1}"

    # ---------------------------------------------------------- SEM CUSTO
    # Os 45 que não formam preço em cenário nenhum. Sai como formulário — uma linha por SKU,
    # com a coluna de custo em branco para preencher e devolver — porque a pergunta ao
    # fornecedor é essa lista, não um relatório sobre ela.
    sc = wb.create_sheet("SEM CUSTO")
    sc.sheet_view.showGridLines = False
    sc.merge_cells("A1:G1")
    h = sc["A1"]
    h.value = ("SKUs SEM CUSTO — preencher e devolver. Sem custo não existe preço em cenário "
               "nenhum, e não há o que derivar.")
    h.font = Font(name=F, bold=True, size=11, color=COTTON)
    h.fill = fill_capa
    h.alignment = Alignment(horizontal="left", vertical="center")
    sc.row_dimensions[1].height = 28
    cabecalho(sc, 2, ["Fornecedor", "Família", "Produto", "Especificação", "Fios",
                      "CUSTO (preencher)", "Observação de quem preencheu"],
              [24, 20, 42, 40, 8, 18, 40])
    prod_por_sku = {p.sku_key: p for p in ctx["produtos"]}
    linha = 3
    sem_custo = [p for p in pendencias if p["tipo"] == "SEM CUSTO"]
    for p in sorted(sem_custo, key=lambda x: (x["fornecedor"], x["familia"] or "", x["nome"])):
        prod = prod_por_sku.get(p["sku"])
        sc.cell(row=linha, column=1, value=p["fornecedor"])
        sc.cell(row=linha, column=2, value=p["familia"])
        sc.cell(row=linha, column=3, value=p["nome"])
        sc.cell(row=linha, column=4, value=(prod.especificacao if prod else "") or "")
        sc.cell(row=linha, column=5, value=getattr(prod, "thread_count", None) if prod else None)
        cel = sc.cell(row=linha, column=6)
        cel.fill = fill_azul
        cel.border = borda_entrada
        cel.number_format = BRL
        for col in range(1, 8):
            sc.cell(row=linha, column=col).border = borda
        linha += 1
    sc.freeze_panes = "A3"
    sc.auto_filter.ref = f"A2:G{linha - 1}"

    # -------------------------------------------------------------- FISCAL
    fs = wb.create_sheet("FISCAL")
    fs.sheet_view.showGridLines = False
    fs.merge_cells("A1:M1")
    h = fs["A1"]
    h.value = ("CARGA FISCAL POR DESTINO — o que está dentro do preço e o que o cliente "
               "paga por fora")
    h.font = Font(name=F, bold=True, size=11, color=COTTON)
    h.fill = fill_capa
    h.alignment = Alignment(horizontal="left", vertical="center")
    fs.row_dimensions[1].height = 28
    fs.merge_cells("A2:M2")
    fs["A2"] = ("A alíquota muda com a NATUREZA da mercadoria: KTC é importada (4% "
                "interestadual), Daune e Decor são nacionais (7% ou 12%). Nas linhas de NÃO "
                "CONTRIBUINTE marcadas em vermelho, a base interna e a incidência do FCP são "
                "SUPOSIÇÃO, não fonte — veja a aba PREMISSAS antes de usar.")
    fs["A2"].font = Font(name=F, italic=True, size=9, color=CINZA)
    fs["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    fs.row_dimensions[2].height = 26
    cabecalho(fs, 3, ["UF", "Destino", "Cliente", "Natureza", "ICMS no preço",
                      "Interestadual", "DIFAL", "FCP", "Quem recolhe o DIFAL",
                      "Reduz margem Anara?", "Interna do destino", "Base interna", "Regra"],
              [6, 20, 17, 13, 13, 13, 11, 9, 20, 18, 15, 13, 80])
    derivadas_uf = {d["uf"] for d in ctx.get("derivados", []) if d["fonte"] == "DERIVADO"}
    linha = 4
    for d in ctx["fiscal_detalhe"]:
        vals = [d["uf"], d["estado"],
                "contribuinte" if d["contribuinte"] else "não contribuinte",
                d["natureza"], d["icms_total"], d["interestadual"], d["difal"], d["fcp"],
                d["quem_recolhe"], d["na_margem"], d["interna_destino"], d["base_interna"],
                d["regra"]]
        for i, v in enumerate(vals, start=1):
            c = fs.cell(row=linha, column=i, value=v)
            if i in (5, 6, 7, 8, 11, 12):
                c.number_format = PCT
            if i == 13:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.font = Font(name=F, size=8, color=CINZA)
        if d["uf"] in derivadas_uf and not d["contribuinte"]:
            fs.cell(row=linha, column=12).fill = fill_verm
        linha += 1
    fs.freeze_panes = "A4"
    fs.auto_filter.ref = f"A3:M{linha - 1}"

    # --------------------------------------------------------------- FRETE
    fr = wb.create_sheet("FRETE")
    fr.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [26, 22, 14, 14, 14, 14, 16, 60]):
        fr.column_dimensions[col].width = w
    fr.merge_cells("A1:H1")
    h = fr["A1"]
    h.value = "FRETE — TABELA TRANSAL. NÃO ESTÁ NO PREÇO UNITÁRIO."
    h.font = Font(name=F, bold=True, size=12, color=COTTON)
    h.fill = PatternFill("solid", fgColor="FF8B2E2E")
    h.alignment = Alignment(horizontal="center", vertical="center")
    fr.row_dimensions[1].height = 28
    fr.merge_cells("A2:H3")
    fr["A2"] = ("A tarifa é por TONELADA, com mínimo por EMBARQUE e adicionais sobre o valor "
                "da NF. Por isso não existe 'frete por peça': o frete unitário de um pedido "
                "de 10 peças e o de um de 10.000 diferem em ordem de grandeza. Use esta aba "
                "para calcular o frete DO PEDIDO e some por fora — ou monte a cotação no "
                "sistema, que faz o rateio por item.")
    fr["A2"].font = Font(name=F, size=10, color=MIDNIGHT)
    fr["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    fr.row_dimensions[2].height = 22

    linha = 5
    for t in ctx["frete"]["tabelas"]:
        cel = fr.cell(row=linha, column=1, value="TABELA")
        cel.font = Font(name=F, bold=True, size=11, color=COTTON)
        cel.fill = fill_secao
        fr.cell(row=linha, column=2,
                value=f"{t.origem_logistica_cidade}/{t.origem_logistica_uf}")
        fr.cell(row=linha, column=5, value=f"tarifa em {t.tarifa_unidade}")
        fr.cell(row=linha, column=7, value=f"cubagem {t.fator_cubagem_kg_m3} kg/m³")
        fr.cell(row=linha, column=8, value=t.documento_fonte)
        linha += 1

        # A origem é da KTC. Itajaí é o ponto de ENTRADA da mercadoria importada — não prova
        # de onde Daune e Decor embarcam, que é a pergunta Q-L, aberta. Sem esta linha,
        # quem consultar calcula frete de travesseiro Daune saindo de Itajaí.
        cel = fr.cell(row=linha, column=1, value="Origem vale para")
        cel.font = Font(name=F, bold=True, size=10, color=MIDNIGHT)
        fr.cell(row=linha, column=2, value="SOMENTE KTC (importada)").fill = fill_bege
        c = fr.cell(row=linha, column=8, value=(
            "Itajaí/SC é o ponto de entrada da mercadoria importada da KTC. De onde Daune e "
            "Decor Tricot embarcam NÃO está determinado (Q-L): para esses dois fornecedores o "
            "frete é A_COTAR — não use esta tabela."))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, size=8, color=CINZA)
        fr.row_dimensions[linha].height = 34
        linha += 1

        # Vigência: a tabela vence, e planilha solta não bloqueia sozinha como o sistema faz.
        cel = fr.cell(row=linha, column=1, value="VÁLIDA ATÉ")
        cel.font = Font(name=F, bold=True, size=11, color=COTTON)
        cel.fill = PatternFill("solid", fgColor="FF8B2E2E")
        cel = fr.cell(row=linha, column=2, value=str(t.valid_to))
        cel.font = Font(name=F, bold=True, size=11, color="FF8B2E2E")
        cel.fill = fill_verm
        c = fr.cell(row=linha, column=8, value=(
            f"Vigência {t.valid_from} a {t.valid_to}. Depois desta data a tarifa está VENCIDA "
            "— não cote com ela. A tabela ainda traz cláusulas de revisão por política de "
            "combustível e por queda de volumetria, que podem antecipar o reajuste "
            "(C-NEW-07)."))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, size=8, color=CINZA)
        fr.row_dimensions[linha].height = 34
        linha += 1

        cel = fr.cell(row=linha, column=1, value="ICMS do frete")
        cel.font = Font(name=F, bold=True, size=10, color=MIDNIGHT)
        fr.cell(row=linha, column=2, value=t.icms_situacao).fill = fill_verm
        c = fr.cell(row=linha, column=8, value=t.icms_notas)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, size=8, color=CINZA)
        fr.row_dimensions[linha].height = 46
        linha += 1

        # A tarifa do pedágio está confirmada; a BASE dela não. São coisas diferentes, e a
        # diferença só aparece quando peso real e peso taxado divergem — ou seja, na cubagem.
        cel = fr.cell(row=linha, column=1, value="Base do pedágio")
        cel.font = Font(name=F, bold=True, size=10, color=MIDNIGHT)
        fr.cell(row=linha, column=2, value=t.pedagio_base).fill = fill_verm
        c = fr.cell(row=linha, column=8, value=(
            "A tarifa de R$ 0,0536/kg está confirmada pelo exemplo da própria planilha, mas a "
            "BASE não: não se sabe se incide sobre peso real ou peso taxado. No exemplo da "
            "TRANSAL os dois valem 500 kg, então ele não decide. Importa quando a cubagem faz "
            "os dois divergirem (C-NEW-08)."))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, size=8, color=CINZA)
        fr.row_dimensions[linha].height = 40
        linha += 2

    fr.cell(row=linha, column=1, value="ADICIONAIS").font = Font(name=F, bold=True, size=11,
                                                                color=MIDNIGHT)
    linha += 1
    cabecalho(fr, linha, ["Código", "Nome", "Tipo", "Valor", "Unidade", "Situação", "",
                          "Regra"])
    linha += 1
    for c_ in ctx["frete"]["componentes"]:
        fr.cell(row=linha, column=1, value=c_.codigo)
        fr.cell(row=linha, column=2, value=c_.nome)
        fr.cell(row=linha, column=3, value=c_.tipo)
        fr.cell(row=linha, column=4, value=c_.valor)
        fr.cell(row=linha, column=5, value=c_.unidade)
        cel = fr.cell(row=linha, column=6, value=c_.situacao)
        if (c_.situacao or "").upper() != "APLICA":
            cel.fill = fill_verm
        cel2 = fr.cell(row=linha, column=8, value=c_.regra)
        cel2.alignment = Alignment(wrap_text=True, vertical="top")
        cel2.font = Font(name=F, size=8, color=CINZA)
        linha += 1

    linha += 1
    fr.cell(row=linha, column=1, value="FAIXAS DE PESO POR REGIÃO").font = \
        Font(name=F, bold=True, size=11, color=MIDNIGHT)
    linha += 1
    cabecalho(fr, linha, ["Região de destino", "Peso de (kg)", "Peso até (kg)",
                          "Tarifa (R$/t)", "Mínimo (R$)", "Prazo", "", "Observação"])
    linha += 1
    for f_ in sorted(ctx["frete"]["faixas"], key=lambda x: (x.regiao_destino, x.peso_de)):
        # As duas tabelas da TRANSAL grafam a mesma região com e sem acento ("ITAJAÍ" na
        # cobertura, "ITAJAI" nas faixas). O motor normaliza e casa certo; a planilha, lida
        # por gente, mostraria dois nomes diferentes para a mesma região e a busca falharia
        # em 110 das 238 cidades. Aqui as duas pontas usam a MESMA chave normalizada.
        fr.cell(row=linha, column=1, value=chave_regiao(f_.regiao_destino))
        fr.cell(row=linha, column=2, value=f_.peso_de)
        fr.cell(row=linha, column=3, value=f_.peso_ate)
        c = fr.cell(row=linha, column=4, value=f_.tarifa)
        c.number_format = BRL
        c2 = fr.cell(row=linha, column=5, value=f_.frete_minimo)
        c2.number_format = BRL
        fr.cell(row=linha, column=6, value=f_.prazo)
        if f_.tarifa is None or f_.frete_minimo is None:
            # Região listada sem tarifa não é frete barato: é frete que não existe na tabela.
            for col in range(1, 7):
                fr.cell(row=linha, column=col).fill = fill_verm
            cel = fr.cell(row=linha, column=8, value=(
                "SEM TARIFA, MÍNIMO OU PRAZO na tabela. Destino nesta região é "
                "FRETE_A_COTAR — não aproxime pela região vizinha."))
            cel.font = Font(name=F, size=8, bold=True, color="FF8B2E2E")
            cel.alignment = Alignment(wrap_text=True, vertical="top")
        linha += 1

    linha += 1
    fr.cell(row=linha, column=1, value="CIDADES ATENDIDAS").font = Font(name=F, bold=True,
                                                                       size=11, color=MIDNIGHT)
    cel = fr.cell(row=linha, column=3, value=(
        "Esta lista é a cobertura INTEIRA. Cidade que não está aqui é FRETE_A_COTAR — "
        "não aproxime por cidade ou região vizinha."))
    cel.font = Font(name=F, bold=True, size=9, color="FF8B2E2E")
    linha += 1
    inicio_cob = linha
    cabecalho(fr, linha, ["Cidade", "UF", "Região de destino", "Unidade"])
    linha += 1
    regioes_com_tarifa = {chave_regiao(f_.regiao_destino)
                          for f_ in ctx["frete"]["faixas"] if f_.tarifa is not None}
    for cb in sorted(ctx["frete"]["cobertura"], key=lambda x: (x.uf, x.cidade)):
        fr.cell(row=linha, column=1, value=cb.cidade)
        fr.cell(row=linha, column=2, value=cb.uf)
        reg = chave_regiao(cb.regiao_destino)
        cel = fr.cell(row=linha, column=3, value=reg)
        if reg not in regioes_com_tarifa:
            cel.fill = fill_verm
        fr.cell(row=linha, column=4, value=cb.unidade)
        linha += 1
    fr.auto_filter.ref = f"A{inicio_cob}:D{linha - 1}"

    # ----------------------------------------------------------- PREMISSAS
    pm = wb.create_sheet("PREMISSAS")
    pm.sheet_view.showGridLines = False
    pm.column_dimensions["A"].width = 34
    pm.column_dimensions["B"].width = 30
    pm.column_dimensions["C"].width = 26
    pm.merge_cells("A1:C1")
    h = pm["A1"]
    h.value = "DE ONDE VEIO CADA NÚMERO"
    h.font = Font(name=F, bold=True, size=12, color=COTTON)
    h.fill = fill_capa
    h.alignment = Alignment(horizontal="center", vertical="center")
    pm.row_dimensions[1].height = 28

    linhas = [("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M"), ""),
              ("Origem fiscal da operação", ctx["origem_nome"], "config catalogo_origem"),
              ("Finalidade padrão", ctx["finalidade"], "config fiscal_finalidade_padrao"),
              ("SKUs ativos", len(ctx["produtos"]), ""),
              ("SKUs com preço formado", len(skus_com_preco), ""),
              ("Cenários liberados", len(ctx["validos"]), f"de {len(cenarios)} possíveis"),
              ("Condições de pagamento", ", ".join(cods), "só encargo confirmado"),
              ("", "", "")]
    for chave, dado in ctx["premissas"].items():
        if isinstance(dado, dict):
            linhas.append((chave, dado.get("valor"),
                           f"vigente desde {dado.get('valid_from')}"))
    for cond in condicoes:
        linhas.append((f"Encargo — {cond.codigo}", para_float(cond.encargo_pct),
                       "confirmado"))

    linha = 2
    for a, b, c in linhas:
        ca = pm.cell(row=linha, column=1, value=a)
        ca.font = Font(name=F, bold=True, size=10, color=MIDNIGHT)
        pm.cell(row=linha, column=2, value=b).font = Font(name=F, size=10)
        cc = pm.cell(row=linha, column=3, value=c)
        cc.font = Font(name=F, size=9, italic=True, color=CINZA)
        if str(a).startswith("Encargo"):
            pm.cell(row=linha, column=2).number_format = PCT
        linha += 1

    # --- o que é dado cadastrado e o que é inferência -----------------------------------
    linha += 1
    pm.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=5)
    cel = pm.cell(row=linha, column=1,
                  value="BASE INTERNA E FCP — o que é fonte e o que é suposição")
    cel.font = Font(name=F, bold=True, size=11, color=COTTON)
    cel.fill = fill_secao
    linha += 1
    pm.merge_cells(start_row=linha, start_column=1, end_row=linha + 3, end_column=5)
    cel = pm.cell(row=linha, column=1, value=(
        "Nos 26 estados marcados abaixo, esta planilha ASSUME QUE O FCP SE APLICA aos "
        "produtos da Anara, na alíquota da coluna FEM — e que não se aplica onde FEM = 0. "
        "Isso não foi verificado. O FCP não incide sobre tudo que entra num estado: a "
        "incidência depende do PRODUTO e muda por UF e por vigência, e é por isso que o "
        "sistema exige regra cadastrada em vez de ler uma coluna por estado.\n\n"
        "A pergunta para o contador é essa — \"o FCP incide sobre roupa de cama e banho "
        "neste estado, e em que alíquota?\" —, e não apenas qual é a base interna. "
        "Enquanto não houver resposta, todo preço a NÃO CONTRIBUINTE fora de SP e RJ "
        "carrega essa suposição."))
    cel.font = Font(name=F, size=9, color=MIDNIGHT)
    cel.alignment = Alignment(wrap_text=True, vertical="top")
    linha += 4
    pm.column_dimensions["D"].width = 14
    pm.column_dimensions["E"].width = 26
    cabecalho(pm, linha, ["UF", "Interna cadastrada", "FEM", "Base interna",
                          "Base e FCP vêm de"])
    linha += 1
    for d in sorted(ctx.get("derivados", []), key=lambda x: x["uf"]):
        pm.cell(row=linha, column=1, value=d["uf"])
        pm.cell(row=linha, column=2, value=d["interna"]).number_format = PCT
        pm.cell(row=linha, column=3, value=d["fem"]).number_format = PCT
        pm.cell(row=linha, column=4, value=d["base"]).number_format = PCT
        if d["fonte"] == "DERIVADO":
            cel = pm.cell(row=linha, column=5, value="SUPOSIÇÃO — confirmar incidência")
            cel.fill = fill_verm
            cel.font = Font(name=F, bold=True, size=9, color="FF8B2E2E")
        else:
            pm.cell(row=linha, column=5, value="CADASTRADO — regra de FCP própria")
        linha += 1

    wb.save(SAIDA)
    return fim_grade, len(skus_com_preco)


def main():
    engine = create_engine(f"sqlite:///file:{os.path.abspath(DB)}?mode=ro&uri=true")
    with Session(engine) as s:
        derivados = instalar_derivacao(s)
        grade, pendencias, interno, cenarios, ctx = calcular(s)
        ctx["derivados"] = derivados
        n_linhas, n_skus = escrever(grade, pendencias, interno, cenarios, ctx)

    validos = len(ctx["validos"])
    print(f"Arquivo:            {SAIDA}")
    print(f"SKUs com preço:     {n_skus} de {len(ctx['produtos'])} ativos")
    print(f"Cenários liberados: {validos} de {len(cenarios)} "
          f"({len(ctx['estados'])} estados × contribuinte/não)")
    print(f"Linhas na GRADE:    {n_linhas:,}".replace(",", "."))
    print(f"Pendências:         {len(pendencias)}")
    if ctx["motivos"]:
        print("\nMotivos de bloqueio de cenário:")
        for m, n in ctx["motivos"].most_common():
            print(f"  {n:3}×  {m}")


if __name__ == "__main__":
    main()
