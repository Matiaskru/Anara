#!/usr/bin/env python3
"""Calculadora KTC em Excel — uma tela com o que se escolhe e o preço. Nada mais.

Duas abas à vista: CALCULADORA e COMO USAR. Toda a conta mora numa aba oculta (MOTOR), junto
com as tabelas de apoio. Quem usa não vê custo, nem tecido por m², nem imposto de importação:
escolhe nas listas e lê o preço.

Decisões que evitam os erros da versão anterior:

* os endereços das células são gerados por contador e guardados num dicionário — nenhuma
  fórmula tem número de linha escrito à mão, então mexer no layout não quebra nada;
* a aba visível não faz conta nenhuma: só lê valores prontos do MOTOR. Assim não existe
  #VALOR! na tela;
* o MOTOR sempre devolve número (0 quando não se aplica) e uma bandeira "dá para calcular",
  em vez de texto no meio de conta;
* nenhuma função sensível a idioma (TEXT com máscara) na planilha.

Rodar:  python3 scripts/gerar_calculadora_excel.py
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.expanduser("~/Anara-Cotacao"))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = os.path.expanduser("~/Downloads/Calculadora KTC Anara.xlsx")
DADOS = json.load(open("/tmp/dados_calculadora.json"))

MIDNIGHT, COPPER, SAND = "FF30354F", "FFC49281", "FFD6C8BE"
COTTON, CINZA = "FFF9F6F2", "FF6B6E7C"
AZUL, VERDE, BEGE = "FFD8E4F5", "FFE3EFE7", "FFF4EFEA"
F = "Calibri"

fill_capa = PatternFill("solid", fgColor=MIDNIGHT)
fill_secao = PatternFill("solid", fgColor=COPPER)
fill_azul = PatternFill("solid", fgColor=AZUL)
fill_verde = PatternFill("solid", fgColor=VERDE)
fill_bege = PatternFill("solid", fgColor=BEGE)
fill_cab = PatternFill("solid", fgColor=MIDNIGHT)
lado = Side(style="thin", color=SAND)
borda = Border(left=lado, right=lado, top=lado, bottom=lado)
grosso = Side(style="medium", color=COPPER)
borda_entrada = Border(left=grosso, right=grosso, top=grosso, bottom=grosso)

BRL = '"R$ "#,##0.00'
PCT = "0.0%"
PCT2 = "0.00%"

wb = Workbook()

# ===========================================================================
# ABAS DE APOIO (ocultas)
# ===========================================================================
def tabela(ws, titulo, cabecalhos, linhas, larguras):
    ws.cell(1, 1, titulo).font = Font(name=F, size=13, bold=True, color=MIDNIGHT)
    for i, texto in enumerate(cabecalhos, start=1):
        c = ws.cell(3, i, texto)
        c.font = Font(name=F, size=10, bold=True, color=COTTON)
        c.fill = fill_cab
    for j, linha in enumerate(linhas, start=4):
        for i, valor in enumerate(linha, start=1):
            c = ws.cell(j, i, valor)
            c.font = Font(name=F, size=10)
            c.border = borda
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    return 4, 3 + len(linhas)


ws = wb.active
ws.title = "TECIDOS"
linhas = [[f"{t['material']} · {'listrado' if t['plain_or_stripe']=='stripe' else 'liso'}",
           t["thread_count"], "COTTON" if (t["cotton_pct"] or 0) >= 0.999 else "CVC",
           t["price_usd_m2"]] for t in DADOS["tecidos"]]
TEC_INI, TEC_FIM = tabela(ws, "TECIDOS KTC", ["Tecido", "Fios", "Composição", "US$/m²"],
                          linhas, [40, 8, 13, 12])

ws = wb.create_sheet("TOALHAS_KG")
linhas = [[f"{t['subcategoria']} · {'listrado' if t['plain_or_stripe']=='stripe' else 'liso'}",
           t["price_usd_kg"]] for t in DADOS["toalhas"]]
TOA_INI, TOA_FIM = tabela(ws, "TOALHAS", ["Construção", "US$/kg"], linhas, [34, 12])

peso_m2 = {p["escopo"]: p["valor"] for p in DADOS["parametros"] if p["chave"] == "peso_kg_m2_familia"}
cmt = {}
for c in DADOS["cmt"]:
    cmt.setdefault(c["familia"], c["cmt_usd"])
ii_fam = {n["familia"]: n for n in DADOS["ncm"]}
FAMILIAS = [
    ("Lençol de cima (plano)", "TECIDO", "Flat Sheet", cmt.get("Flat Sheet", .75), 1, 4, 4,
     peso_m2.get("Flat Sheet"), "LENCOL"),
    ("Capa de duvet", "TECIDO", "Duvet Cover", cmt.get("Duvet Cover", 1.5), 2, 4, 4,
     peso_m2.get("Duvet Cover"), "OUTRO"),
    ("Toalha de banho", "TOALHA", "Bath Towel", 0, 0, 0, 0, 0, "TOALHA"),
    ("Toalha de rosto", "TOALHA", "Hand Towel", 0, 0, 0, 0, 0, "TOALHA"),
    ("Toalha de piso", "TOALHA", "Bath Mat", 0, 0, 0, 0, 0, "TOALHA"),
    ("Toalha de piscina", "TOALHA", "Pool Towel", 0, 0, 0, 0, 0, "TOALHA"),
    ("Toalha de lavabo", "TOALHA", "Wash Cloth", 0, 0, 0, 0, 0, "TOALHA"),
    ("Lençol com elástico", "PEDIR", "Fitted Sheet", 0, 0, 0, 0, 0, "LENCOL"),
    ("Fronha", "PEDIR", "Pillow Case", 0, 0, 0, 0, 0, "OUTRO"),
    ("Roupão", "PEDIR", "Bathrobe", 0, 0, 0, 0, 0, "TOALHA"),
    ("Edredom / insert", "PEDIR", "Duvet Insert", 0, 0, 0, 0, 0, "OUTRO"),
    ("Protetor de colchão", "PEDIR", "Mattress Protector", 0, 0, 0, 0, 0, "OUTRO"),
    ("Topper de colchão", "PEDIR", "Mattress Topper", 0, 0, 0, 0, 0, "OUTRO"),
    ("Chinelo", "PEDIR", "Slipper", 0, 0, 0, 0, 0, "OUTRO"),
]
ws = wb.create_sheet("FAMILIAS")
linhas = [[r, t, n, c or 0, p or 0, hw or 0, hl or 0, pm or 0, g,
           (ii_fam.get(n, {}).get("ii_preferencial") or 0)]
          for r, t, n, c, p, hw, hl, pm, g in FAMILIAS]
FAM_INI, FAM_FIM = tabela(ws, "FAMÍLIAS",
                          ["Produto", "Tipo", "Técnico", "CMT", "Painéis", "Bainha L",
                           "Bainha C", "Peso kg/m²", "Grupo", "I.I."], linhas,
                          [26, 10, 20, 9, 9, 10, 10, 11, 14, 9])

p = DADOS["premissas"]
PARAMS = [("fx", "Câmbio USD/BRL", p.get("fx_usd_brl", 5.11)),
          ("frete", "Frete internacional US$/kg", p.get("frete_int_usd_kg", .516)),
          ("desp", "Outras despesas US$/un", p.get("outras_desp_usd_un", .2487532709)),
          ("pis", "PIS/COFINS", p.get("pis_cofins_pct", .0759)),
          ("shr_cvc", "Encolhimento CVC", .03), ("shr_cot", "Encolhimento algodão", .05),
          ("waste", "Waste", .03), ("qual", "Perda de 2ª qualidade", .01),
          ("mktc", "Margem da KTC", .15)]
ws = wb.create_sheet("PARAMETROS")
PAR_INI, PAR_FIM = tabela(ws, "PARÂMETROS", ["Chave", "Parâmetro", "Valor"],
                          [[c, r, v] for c, r, v in PARAMS], [12, 40, 16])
PAR = {c: f"PARAMETROS!$C${PAR_INI + i}" for i, (c, _r, _v) in enumerate(PARAMS)}

ws = wb.create_sheet("FISCAL")
linhas = [[e["estado"], e["aliquota_interestadual"], e["aliquota_interna"], e["carga_final"]]
          for e in DADOS["estados"]]
EST_INI, EST_FIM = tabela(ws, "ESTADOS", ["Estado", "Interestadual", "Interna", "Carga final"],
                          linhas, [24, 14, 14, 14])

ws = wb.create_sheet("PAGAMENTO")
linhas = [[c["label"], c["encargo_pct"] or 0] for c in DADOS["pagamento"]]
PAG_INI, PAG_FIM = tabela(ws, "PAGAMENTO", ["Condição", "Encargo"], linhas, [30, 12])

ws = wb.create_sheet("MARGENS")
linhas = [["TOALHA", .12], ["LENCOL_MENOR", .16], ["LENCOL_MAIOR", .18], ["OUTRO", .15]]
MAR_INI, MAR_FIM = tabela(ws, "MARGENS PADRÃO", ["Grupo", "Margem"], linhas, [18, 12])
MAR = {ws.cell(l, 1).value: f"MARGENS!$B${l}" for l in range(MAR_INI, MAR_FIM + 1)}

ws = wb.create_sheet("COMISSAO")
faixas = DADOS["comissao"]
linhas = [[f[0], (faixas[i + 1][0] if i + 1 < len(faixas) else 99), f[1]]
          for i, f in enumerate(faixas)]
COM_INI, COM_FIM = tabela(ws, "COMISSÃO", ["Markup mín.", "Markup máx.", "Comissão"], linhas,
                          [14, 14, 12])

MEDIDAS = ["30x30", "45x70", "45x75", "45x80", "45x85", "48x80", "48x85", "50x70", "50x80",
           "50x85", "50x90", "50x100", "55x75", "58x150", "60x190", "60x240", "60x270", "60x285",
           "65x65", "70x135", "70x140", "70x150", "86x150", "88x188", "90x150", "90x160",
           "90x170", "100x150", "100x200", "120x280", "140x200", "150x220", "156x230", "160x200",
           "160x230", "160x310", "180x200", "180x240", "180x280", "180x295", "190x250", "190x260",
           "190x290", "200x200", "200x290", "210x220", "220x230", "220x240", "230x220", "240x230",
           "240x250", "240x260", "240x280", "250x260", "250x290", "260x250", "260x280", "260x290",
           "270x265", "270x275", "280x290", "285x265", "290x260", "300x300", "325x295"]
GRAMATURAS = [330, 380, 400, 420, 450, 500, 550, 600, 650, 670, 750, 950]
QUANTIDADES = [1, 10, 20, 30, 50, 100, 150, 200, 300, 500, 750, 1000, 2000]
ws = wb.create_sheet("LISTAS")
for i, m in enumerate(MEDIDAS):
    ws.cell(1 + i, 1, m)
for i, g in enumerate(GRAMATURAS):
    ws.cell(1 + i, 2, g)
for i, q in enumerate(QUANTIDADES):
    ws.cell(1 + i, 3, q)
for i, s in enumerate(["SIM", "NÃO"]):
    ws.cell(1 + i, 4, s)
for i, s in enumerate(["liso", "listrado"]):
    ws.cell(1 + i, 5, s)

for n, ref in [("LISTA_PRODUTOS", f"FAMILIAS!$A${FAM_INI}:$A${FAM_FIM}"),
               ("LISTA_TECIDOS", f"TECIDOS!$A${TEC_INI}:$A${TEC_FIM}"),
               ("LISTA_ESTADOS", f"FISCAL!$A${EST_INI}:$A${EST_FIM}"),
               ("LISTA_PAGAMENTO", f"PAGAMENTO!$A${PAG_INI}:$A${PAG_FIM}"),
               ("LISTA_MEDIDAS", f"LISTAS!$A$1:$A${len(MEDIDAS)}"),
               ("LISTA_GRAMATURAS", f"LISTAS!$B$1:$B${len(GRAMATURAS)}"),
               ("LISTA_QUANTIDADES", f"LISTAS!$C$1:$C${len(QUANTIDADES)}"),
               ("LISTA_SIMNAO", "LISTAS!$D$1:$D$2"),
               ("LISTA_ACABAMENTO", "LISTAS!$E$1:$E$2")]:
    wb.defined_names.add(DefinedName(n, attr_text=ref))

# ===========================================================================
# CALCULADORA — só o que se escolhe e o preço
# ===========================================================================
ws = wb.create_sheet("CALCULADORA", 0)
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 130
for coluna, largura in [("A", 3), ("B", 34), ("C", 30), ("D", 3)]:
    ws.column_dimensions[coluna].width = largura

E = {}          # nome do campo → endereço da célula. Única fonte de verdade.
linha = 1


def espaco(altura=8):
    global linha
    ws.row_dimensions[linha].height = altura
    linha += 1


def capa(texto, subtitulo):
    global linha
    ws.merge_cells(start_row=linha, start_column=2, end_row=linha + 1, end_column=3)
    c = ws.cell(linha, 2, texto)
    c.font = Font(name=F, size=19, bold=True, color=COTTON)
    c.fill = fill_capa
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(linha, 3).fill = fill_capa
    ws.row_dimensions[linha].height = 30
    ws.row_dimensions[linha + 1].height = 14
    linha += 2
    ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=3)
    c = ws.cell(linha, 2, subtitulo)
    c.font = Font(name=F, size=10.5, color=CINZA, italic=True)
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[linha].height = 20
    linha += 1


def secao(texto):
    global linha
    ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=3)
    c = ws.cell(linha, 2, f"   {texto}")
    c.font = Font(name=F, size=11.5, bold=True, color=COTTON)
    c.alignment = Alignment(vertical="center")
    ws.cell(linha, 3).fill = fill_secao
    c.fill = fill_secao
    ws.row_dimensions[linha].height = 24
    linha += 1


def campo(chave, texto, valor=None, formato=None):
    """Uma pergunta com a resposta em célula azul. Guarda o endereço em E."""
    global linha
    ws.cell(linha, 2, texto).font = Font(name=F, size=12, color=MIDNIGHT)
    ws.cell(linha, 2).alignment = Alignment(vertical="center")
    c = ws.cell(linha, 3, valor)
    c.font = Font(name=F, size=12, bold=True, color=MIDNIGHT)
    c.fill = fill_azul
    c.border = borda_entrada
    c.alignment = Alignment(horizontal="center", vertical="center")
    if formato:
        c.number_format = formato
    E[chave] = f"$C${linha}"
    ws.row_dimensions[linha].height = 26
    linha += 1
    return c


espaco()
capa("ANARA · CALCULADORA DE PREÇO", "Escolha nas listas. O preço aparece no fim.")
espaco()
secao("1 · O QUE O CLIENTE PEDIU")
campo("produto", "Produto", "Lençol de cima (plano)")
campo("medida", "Medida (cm)", "240x260")
campo("tecido", "Tecido", "300TC Sateen 100% Cotton · liso")
campo("gramatura", "Gramatura (só toalha)", None)
campo("acabamento", "Liso ou listrado", "liso")
campo("quantidade", "Quantidade de peças", 50, "#,##0")
espaco()
secao("2 · PARA ONDE VAI A VENDA")
campo("origem", "Sai de", "São Paulo")
campo("destino", "Vai para", "São Paulo")
campo("contribuinte", "Cliente é contribuinte de ICMS?", "SIM")
campo("pagamento", "Como o cliente vai pagar", "30 dias")
espaco()
secao("3 · MARGEM  (deixe vazio para usar a padrão)")
campo("margem", "Margem que você quer", None, PCT)
espaco()

secao("PREÇO")
LINHA_PRECO = linha
ws.merge_cells(start_row=linha, start_column=2, end_row=linha + 1, end_column=2)
c = ws.cell(linha, 2, "PREÇO POR PEÇA")
c.font = Font(name=F, size=13, bold=True, color=MIDNIGHT)
c.alignment = Alignment(vertical="center")
ws.merge_cells(start_row=linha, start_column=3, end_row=linha + 1, end_column=3)
c = ws.cell(linha, 3)
c.font = Font(name=F, size=30, bold=True, color=MIDNIGHT)
c.fill = fill_verde
c.number_format = BRL
c.alignment = Alignment(horizontal="center", vertical="center")
E["preco"] = f"$C${linha}"
ws.row_dimensions[linha].height = 34
ws.row_dimensions[linha + 1].height = 16
linha += 2

ws.cell(linha, 2, "TOTAL DA VENDA").font = Font(name=F, size=12, color=MIDNIGHT)
ws.cell(linha, 2).alignment = Alignment(vertical="center")
c = ws.cell(linha, 3)
c.font = Font(name=F, size=15, bold=True, color=MIDNIGHT)
c.fill = fill_bege
c.number_format = BRL
c.alignment = Alignment(horizontal="center", vertical="center")
E["total"] = f"$C${linha}"
ws.row_dimensions[linha].height = 24
linha += 1
espaco()

ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=3)
c = ws.cell(linha, 2)
c.font = Font(name=F, size=11, color="FFA2452A")
c.alignment = Alignment(vertical="center", wrap_text=True)
E["aviso"] = f"$B${linha}"
ws.row_dimensions[linha].height = 34
FIM_TELA = linha

# ===========================================================================
# MOTOR (oculto) — toda a conta, um valor por linha
# ===========================================================================
mot = wb.create_sheet("MOTOR")
mot.column_dimensions["A"].width = 34
mot.column_dimensions["B"].width = 22
mot.cell(1, 1, "MOTOR — a conta inteira. A tela só lê daqui.").font = Font(
    name=F, size=12, bold=True, color=MIDNIGHT)
M = {}
lm = 3


def C(chave):
    """Endereço de um campo da tela, visto de dentro do MOTOR."""
    return f"CALCULADORA!{E[chave]}"


def m(chave, rotulo, formula, formato=None):
    global lm
    mot.cell(lm, 1, rotulo).font = Font(name=F, size=10, color=CINZA)
    c = mot.cell(lm, 2, formula)
    c.font = Font(name=F, size=10, color=MIDNIGHT)
    if formato:
        c.number_format = formato
    M[chave] = f"$B${lm}"
    lm += 1
    return M[chave]


def V(chave):
    return M[chave]


FAM_A = f"FAMILIAS!$A${FAM_INI}:$A${FAM_FIM}"
TEC_A = f"TECIDOS!$A${TEC_INI}:$A${TEC_FIM}"
EST_A = f"FISCAL!$A${EST_INI}:$A${EST_FIM}"
PAG_A = f"PAGAMENTO!$A${PAG_INI}:$A${PAG_FIM}"
TOA_A = f"TOALHAS_KG!$A${TOA_INI}:$A${TOA_FIM}"


def busca_fam(letra, padrao='""'):
    return (f'IFERROR(INDEX(FAMILIAS!${letra}${FAM_INI}:${letra}${FAM_FIM},'
            f'MATCH({C("produto")},{FAM_A},0)),{padrao})')


def busca_tec(letra, padrao="0"):
    return (f'IFERROR(INDEX(TECIDOS!${letra}${TEC_INI}:${letra}${TEC_FIM},'
            f'MATCH({C("tecido")},{TEC_A},0)),{padrao})')


m("tipo", "Tipo de cálculo", f"={busca_fam('B')}")
m("tecnica", "Nome técnico", f"={busca_fam('C')}")
m("larg", "Largura (cm)",
  f'=IFERROR(VALUE(LEFT({C("medida")},FIND("x",{C("medida")})-1)),0)', "0")
m("comp", "Comprimento (cm)",
  f'=IFERROR(VALUE(MID({C("medida")},FIND("x",{C("medida")})+1,10)),0)', "0")
m("gsm", "Gramatura", f'=IFERROR(VALUE({C("gramatura")}),0)', "0")
m("qtd", "Quantidade", f'=IFERROR(VALUE({C("quantidade")}),0)', "0")
m("preco_m2", "Preço do tecido US$/m²", f"={busca_tec('D')}")
m("fios", "Fios", f"={busca_tec('B')}", "0")
m("comp_tec", "Composição", f'={busca_tec("C", chr(34)*2)}')
m("shrink", "Encolhimento",
  f'=IF({V("comp_tec")}="COTTON",{PAR["shr_cot"]},{PAR["shr_cvc"]})', PCT2)
m("cmt", "CMT US$", f"={busca_fam('D', '0')}")
m("paineis", "Painéis", f"={busca_fam('E', '0')}", "0")
m("hw", "Bainha largura", f"={busca_fam('F', '0')}")
m("hl", "Bainha comprimento", f"={busca_fam('G', '0')}")
m("pm2", "Peso kg/m² da família", f"={busca_fam('H', '0')}")
m("grupo", "Grupo de margem", f"={busca_fam('I')}")
m("ii", "Imposto de Importação", f"={busca_fam('J', '0')}", PCT2)

m("exw_tec", "EXW tecido plano US$",
  f'=IF({V("tipo")}<>"TECIDO",0,IFERROR(((({V("larg")}+{V("hw")})*(1+{V("shrink")}))*'
  f'(({V("comp")}+{V("hl")})*(1+{V("shrink")}))/10000*{V("paineis")}/(1-{PAR["waste"]})*'
  f'{V("preco_m2")}+{V("cmt")})/(1-{PAR["qual"]})/(1-{PAR["mktc"]}),0))')
m("peso_toa", "Peso da toalha kg",
  f'=IF({V("tipo")}<>"TOALHA",0,{V("larg")}*{V("comp")}*{V("gsm")}/10000000)')
m("kg_toa", "Preço por kg US$",
  f'=IF({V("tipo")}<>"TOALHA",0,IFERROR(INDEX(TOALHAS_KG!$B${TOA_INI}:$B${TOA_FIM},'
  f'MATCH({V("tecnica")}&" · "&{C("acabamento")},{TOA_A},0)),0))')
m("exw_toa", "EXW toalha US$", f'={V("peso_toa")}*{V("kg_toa")}')
m("exw", "EXW aplicado US$",
  f'=IF({V("tipo")}="TECIDO",{V("exw_tec")},IF({V("tipo")}="TOALHA",{V("exw_toa")},0))')
m("peso", "Peso considerado kg",
  f'=IF({V("tipo")}="TOALHA",{V("peso_toa")},{V("larg")}*{V("comp")}/10000*{V("pm2")})')
m("frete", "Frete internacional US$", f'={V("peso")}*{PAR["frete"]}')
m("net_usd", "Custo NET US$",
  f'={V("exw")}+{V("frete")}+({V("exw")}+{V("frete")})*{V("ii")}+{PAR["desp"]}')
m("net", "Custo NET R$", f'={V("net_usd")}*{PAR["fx"]}', BRL)

m("icms", "ICMS da venda",
  f'=IFERROR(IF({C("origem")}={C("destino")},'
  f'INDEX(FISCAL!$C${EST_INI}:$C${EST_FIM},MATCH({C("destino")},{EST_A},0)),'
  f'IF({C("contribuinte")}="SIM",INDEX(FISCAL!$B${EST_INI}:$B${EST_FIM},'
  f'MATCH({C("destino")},{EST_A},0)),INDEX(FISCAL!$D${EST_INI}:$D${EST_FIM},'
  f'MATCH({C("destino")},{EST_A},0)))),0)', PCT2)
m("pis", "PIS/COFINS", f'={PAR["pis"]}', PCT2)
m("encargo", "Encargo do prazo",
  f'=IFERROR(INDEX(PAGAMENTO!$B${PAG_INI}:$B${PAG_FIM},MATCH({C("pagamento")},{PAG_A},0)),0)',
  PCT2)
m("taxa", "Soma sobre o faturamento", f'={V("icms")}+{V("pis")}+{V("encargo")}', PCT2)
m("m_padrao", "Margem padrão",
  f'=IF({V("grupo")}="TOALHA",{MAR["TOALHA"]},IF({V("grupo")}="LENCOL",'
  f'IF({V("fios")}<300,{MAR["LENCOL_MENOR"]},{MAR["LENCOL_MAIOR"]}),{MAR["OUTRO"]}))', PCT2)
m("margem", "Margem usada",
  f'=IF({C("margem")}="",{V("m_padrao")},{C("margem")})', PCT2)

# faixas de comissão
mot.cell(lm, 1, "Faixas (mín · máx · comissão · markup · vale?)").font = Font(
    name=F, size=9, color=CINZA)
lm += 1
faixa_ini = lm
for i in range(len(faixas)):
    origem = COM_INI + i
    mot.cell(lm, 1, f"=COMISSAO!$A${origem}").number_format = "0%"
    mot.cell(lm, 2, f"=COMISSAO!$B${origem}").number_format = "0%"
    mot.cell(lm, 3, f"=COMISSAO!$C${origem}").number_format = "0%"
    mot.cell(lm, 4, f'=IFERROR({V("margem")}/(1-{V("taxa")}-$C${lm}-{V("margem")}),-99)'
             ).number_format = PCT2
    mot.cell(lm, 5, f'=IF(AND($D${lm}>=$A${lm},$D${lm}<$B${lm}),1,0)')
    lm += 1
faixa_fim = lm - 1
lm += 1
m("comissao", "Comissão resolvida",
  f'=IFERROR(INDEX($C${faixa_ini}:$C${faixa_fim},MATCH(1,$E${faixa_ini}:$E${faixa_fim},0)),'
  f'$C${faixa_fim})', PCT2)
m("markup", "Markup implícito",
  f'=IFERROR(INDEX($D${faixa_ini}:$D${faixa_fim},MATCH(1,$E${faixa_ini}:$E${faixa_fim},0)),'
  f'$D${faixa_fim})', PCT2)
m("preco", "PREÇO POR PEÇA R$",
  f'=IFERROR({V("net")}*(1+{V("markup")})/(1-{V("taxa")}-{V("comissao")}),0)', BRL)
m("calculavel", "Dá para calcular? (1/0)",
  f'=IF({V("tipo")}="TECIDO",IF(AND({V("preco_m2")}>0,{V("larg")}>0),1,0),'
  f'IF({V("tipo")}="TOALHA",IF(AND({V("kg_toa")}>0,{V("gsm")}>0,{V("larg")}>0),1,0),0))', "0")
m("situacao", "Aviso para a tela",
  f'=IF({V("tipo")}="PEDIR",'
  f'"Este produto a KTC nunca explicou como calcula. Peça uma cotação a eles.",'
  f'IF({V("calculavel")}=1,"",'
  f'IF({V("tipo")}="TECIDO","Escolha o tecido e a medida para o preço aparecer.",'
  f'IF(AND({V("tipo")}="TOALHA",{V("kg_toa")}=0),'
  f'"Não temos preço para essa toalha. Peça uma cotação à KTC.",'
  f'"Escolha a gramatura e a medida para o preço aparecer."))))')

# --- a tela só lê o motor --------------------------------------------------
ws[E["preco"].replace("$", "")] = f'=IF(MOTOR!{V("calculavel")}=1,MOTOR!{V("preco")},"")'
ws[E["total"].replace("$", "")] = (f'=IF(MOTOR!{V("calculavel")}=1,'
                                   f'MOTOR!{V("preco")}*MOTOR!{V("qtd")},"")')
ws[E["aviso"].replace("$", "")] = f'=MOTOR!{V("situacao")}'

for formula, chave in [("=LISTA_PRODUTOS", "produto"), ("=LISTA_MEDIDAS", "medida"),
                       ("=LISTA_TECIDOS", "tecido"), ("=LISTA_GRAMATURAS", "gramatura"),
                       ("=LISTA_ACABAMENTO", "acabamento"), ("=LISTA_QUANTIDADES", "quantidade"),
                       ("=LISTA_ESTADOS", "origem"), ("=LISTA_ESTADOS", "destino"),
                       ("=LISTA_SIMNAO", "contribuinte"), ("=LISTA_PAGAMENTO", "pagamento")]:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(ws[E[chave].replace("$", "")])

# ===========================================================================
# COMO USAR
# ===========================================================================
ws = wb.create_sheet("COMO USAR", 1)
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 96
TEXTOS = [
    ("COMO USAR", "titulo"),
    ("", None),
    ("Preencha os campos AZUIS de cima para baixo. Todos são listas: é só escolher.", None),
    ("O preço aparece no fim da tela, em verde.", None),
    ("", None),
    ("O campo de margem pode ficar vazio", "secao"),
    ("Vazio, ele usa a margem padrão daquele produto. Se quiser outra, digite (por exemplo 15%) "
     "e o preço se ajusta.", None),
    ("", None),
    ("O preço já vem completo", "secao"),
    ("O que aparece na tela já tem imposto, o encargo do prazo de pagamento e a comissão dentro. "
     "É o preço para passar ao cliente.", None),
    ("", None),
    ("Se aparecer um aviso em vermelho", "secao"),
    ("Ou falta escolher alguma coisa (o aviso diz o quê), ou é um produto que a fábrica nunca "
     "explicou como calcula. Nesse caso, peça uma cotação a eles.", None),
    ("", None),
    ("Duas coisas mudam o preço e passam despercebidas", "secao"),
    ("Para onde vai a venda: o imposto muda de estado para estado.", None),
    ("Se o cliente é contribuinte de ICMS: isso se pergunta ao cliente, não dá para adivinhar "
     "pelo estado.", None),
    ("", None),
    ("Para mudar preço de tecido, câmbio ou margem padrão", "secao"),
    ("Fale com o Matias. Esses valores ficam em abas ocultas e mexer neles muda o preço de tudo.",
     None),
]
linha = 2
for texto, tipo in TEXTOS:
    c = ws.cell(linha, 2, texto)
    if tipo == "titulo":
        c.font = Font(name=F, size=18, bold=True, color=MIDNIGHT)
    elif tipo == "secao":
        c.font = Font(name=F, size=12, bold=True, color=COPPER)
    else:
        c.font = Font(name=F, size=11, color=MIDNIGHT)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[linha].height = 32 if len(texto) > 90 else 18
    linha += 1

for aba in ("MOTOR", "TECIDOS", "TOALHAS_KG", "FAMILIAS", "PARAMETROS", "FISCAL", "PAGAMENTO",
            "MARGENS", "COMISSAO", "LISTAS"):
    wb[aba].sheet_state = "hidden"
wb.active = 0
wb.save(SAIDA)

# ===========================================================================
# AUDITORIA — o que faltou da última vez
# ===========================================================================
import openpyxl

conf = openpyxl.load_workbook(SAIDA)
tela, motor = conf["CALCULADORA"], conf["MOTOR"]
entradas = {addr.replace("$", "") for chave, addr in E.items()
            if chave not in ("preco", "total", "aviso")}
problemas = []

# 1. toda referência a CALCULADORA feita pelo MOTOR tem de cair num campo de entrada
for linha_m in motor.iter_rows():
    for c in linha_m:
        if isinstance(c.value, str) and "CALCULADORA!" in c.value:
            for ref in re.findall(r"CALCULADORA!\$?([A-Z]+)\$?(\d+)", c.value):
                endereco = f"{ref[0]}{ref[1]}"
                if endereco not in entradas:
                    problemas.append(f"MOTOR!{c.coordinate} aponta para CALCULADORA!{endereco}, "
                                     f"que não é campo de entrada")
                elif tela[endereco].value is None and endereco not in (
                        E["gramatura"].replace("$", ""), E["margem"].replace("$", "")):
                    problemas.append(f"CALCULADORA!{endereco} está vazio")

# 2. a tela não pode fazer conta: só ler o motor
for linha_t in tela.iter_rows():
    for c in linha_t:
        if isinstance(c.value, str) and c.value.startswith("="):
            if "MOTOR!" not in c.value:
                problemas.append(f"CALCULADORA!{c.coordinate} tem fórmula que não vem do MOTOR")

# 3. cada campo de entrada tem lista suspensa
com_lista = set()
for dv in tela.data_validations.dataValidation:
    com_lista.update(str(dv.sqref).split())
for chave, addr in E.items():
    if chave in ("preco", "total", "aviso", "margem"):
        continue
    if addr.replace("$", "") not in com_lista:
        problemas.append(f"campo {chave} ({addr}) ficou sem lista suspensa")

print(f"Planilha gerada: {SAIDA}  ({os.path.getsize(SAIDA):,} bytes)")
print("Abas visíveis:", [a for a in conf.sheetnames if conf[a].sheet_state == "visible"])
print("\nCampos da tela:")
for chave, addr in E.items():
    rotulo = tela[f"B{addr.split('$')[2]}"].value
    print(f"  {chave:14s} {addr:8s} {str(rotulo)[:40]}")
print("\nAUDITORIA:", "tudo certo" if not problemas else f"{len(problemas)} PROBLEMA(S)")
for p in problemas:
    print("  ✗", p)
